#!/usr/bin/env python3
"""Quick mode runner for one-off or batch browser-use quick tasks.

This script stays at the skill layer. It does not change browser-use internals;
it prepares quick queries, optional reference context, and invokes the
browser-use quick/observe CLI commands.
"""
from __future__ import annotations

import argparse
import asyncio
import base64
import csv
import hashlib
import html
import json
import os
import re
import shutil
import shlex
import subprocess
import sys
import time
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

_SCRIPTS_DIR = Path(__file__).resolve().parent
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))
from quick_completion_validate import (
    CompletionValidation,
    validate_quick_attempt,
    write_validation_artifact,
)
from check_ui_test_env import (
    _exe_real_key,
    find_usable_interpreter,
    python_range_label,
)
from skill_config import (
    extract_url,
    preflight_quick,
    resolve_quick_max_retries,
    resolve_quick_max_steps,
    resolve_target_url,
    sanitize_quick_query_for_cli,
)


CASE_COLUMNS = {
    "title": ("用例标题", "标题", "case_title", "title", "name"),
    "precondition": ("前置条件", "前置", "precondition", "preconditions"),
    "steps": ("测试步骤", "步骤", "steps", "test_steps"),
    "expected": ("预期结果", "预期", "expected", "expected_result"),
    "priority": ("等级", "优先级", "priority"),
    "module1": ("一级模块", "模块", "module"),
    "module2": ("二级模块", "子模块", "submodule"),
    "ui_auto": ("是否UI自动化", "是否 UI 自动化", "ui_auto", "ui automation"),
    "description": ("描述", "说明", "description"),
    "case_type": ("类型", "type"),
}

INPUT_KINDS = ("auto", "cases", "requirement", "analysis")
DOC_KEYWORDS_REQUIREMENT = (
    "需求",
    "prd",
    "产品说明",
    "验收标准",
    "用户故事",
    "交互",
    "设计稿",
)
DOC_KEYWORDS_ANALYSIS = (
    "测分",
    "测试分析",
    "测试点",
    "覆盖策略",
    "风险点",
    "功能性测试分析",
)
DOC_TEXT_SUFFIXES = {".md", ".markdown", ".txt", ".csv", ".tsv", ".json", ".yaml", ".yml"}
DEFAULT_DOC_CHUNK_CHARS = 7000
DEFAULT_CASES_PER_DOC_CHUNK = 0
PROMPT_CONTEXT_DIRNAME = "prompt_context"
PERFERENCE_CONTEXT_DIRNAME = "perference_context"
UPLOAD_FILES_DIRNAME = "upload_files"
REGRESSION_DIRNAME = "regression"
CONTEXT_IGNORE_NAMES = {".DS_Store", ".gitkeep"}
DEFAULT_PROMPT_CONTEXT_FILE = "global_readme.md"
UPLOAD_FIXTURE_RULE_MARKER = "上传测试文件选择规则"
UPLOAD_EXPLICIT_PATH_RULE_MARKER = "上传文件路径硬规则"
LOG_QUERY_LIMIT = 1800
LOG_SUMMARY_LIMIT = 900
REPAIR_FOCUS_LIMIT = 1400
LOG_FULL_QUERY_LIMIT = 8000
SMOKE_WALKTHROUGH_MARKERS = (
    "冒烟验证",
    "冒烟测试",
    "前端走查",
    "页面走查",
    "主页面走查",
    "主模块走查",
    "可用性走查",
)
HEALTH_CHECK_MARKERS = (
    "巡检",
    "页面巡检",
    "页面探活",
    "探活",
    "页面存活",
    "可用性巡检",
)
# 每条 quick 首轮均注入的通用 Agent 约束（非业务特性；业务口径见 config/prompt_context/）。
QUICK_AGENT_GUARDRAILS = (
    "下拉框/单选/格式选择：若 DOM 元素是原生 select，优先调用 get_dropdown_options 后"
    "使用 select_dropdown_option 精确选择；只有自定义下拉框才先点触发器展开，再点击目标选项文字。"
    "不要连续 3 次只点触发器本身；若自定义下拉点击后显示值仍未变，"
    "点击下拉获焦后用 ArrowDown 移到目标项再 Enter；仍无效则滚动重定位或结束并报告卡点。\n"
    "指定模型/模式/选项的任务：若用户明确指定模型名、模式名、格式或其它下拉选项，"
    "必须先读取当前显示值；若不是目标值，先完成切换并再次读取页面当前显示值确认。"
    "确认目标值前禁止上传文件、输入业务文本或点击提交/生成；点击提交/生成前还要二次确认当前显示值仍为目标值，"
    "若页面因刷新、上传或切换模式回到默认值，必须先重新切回目标值再继续。\n"
    "生成/提交类任务的本次执行要求：若用户要求输入内容并生成/提交/发送，"
    "必须在本轮操作中实际完成目标值确认、输入用户指定内容、点击提交/生成/发送按钮。"
    "页面上已有的历史记录、旧生成结果、缓存预览或先前会话结果只能作为背景信息，"
    "不能作为本次任务成功依据；即使历史结果的模型和提示词完全匹配，也必须重新执行输入和提交。"
    "若由于权限、余额、开通、缺少上传材料、按钮不可用等原因不能重新提交，必须 done(success=false) 说明原因。\n"
    "上传文件路径：若任务、用例或系统提示中已经给出本地文件路径，upload.path 必须逐字使用该真实路径；"
    "禁止改写、截断、相对化或替换为 /path/to/...、example、sample、image.jpg 等占位路径。"
    "若路径不存在或类型不匹配，必须报告路径问题或从明确列出的内置样例路径中选择，不能臆造新路径。\n"
    "异步/生成类任务：点击提交/生成后，若页面出现生成中、创建中、处理中、任务执行中、AI 正在回复、"
    "排队中、加载中等进行中状态，说明任务已被受理；不要再回头反复修复上传/下拉/输入等前置步骤，"
    "也不要因历史/残留的红色校验提示立即失败。应分段等待并复查，累计至少 60 秒或直到出现明确成功结果、"
    "新结果预览、明确失败状态/错误 toast。进行中状态下发现的前置控件疑点只记为人工核对项，"
    "不能作为中断正在执行任务的理由。\n"
    "反循环：对同一控件连续操作 2 次后，若下拉值/开关/必填红框/toast 警告等关键状态未变，"
    "禁止第 3 次重复同动作；必须换策略（键盘选择、补必填项、滚动重定位、关弹窗重开、切换 Tab 再回来）"
    "或在最终报告中说明卡点；普通 quick 可 done(success=false)，冒烟/巡检必须按对应报告模板结束，不要阶段性失败早停；"
    "提交类按钮（完成/确认）若 2 次点击后页面无跳转/无成功提示/警告仍在，"
    "先检查表单校验或页面警告，不要盲点多按。"
)
SMOKE_ISSUE_KEYWORDS = (
    "报错",
    "错误",
    "异常",
    "红色",
    "红字",
    "红色警告",
    "警告",
    "warning",
    "error",
    "failed",
    "失败",
    "不可用",
    "无法",
    "空白",
    "白屏",
    "崩溃",
    "无响应",
    "加载失败",
    "请求失败",
    "接口错误",
    "权限错误",
    "404",
    "500",
)
SMOKE_NORMAL_RISK_KEYWORDS = (
    "验证码",
    "滑块",
    "人机校验",
    "安全验证",
    "风控",
    "二次校验",
)
TEXT_TIMESTAMP_RULE_MARKER = "文本唯一性规则"
TEXT_INPUT_HINT_PATTERNS = (
    "输入",
    "填写",
    "填入",
    "新建",
    "创建",
    "新增",
    "编辑",
    "保存",
    "名称",
    "标题",
    "描述",
    "备注",
    "问题",
    "答案",
    "标签",
    "库名",
    "用户名",
    "昵称",
)
CDP_INTENT_PATTERNS = (
    "cdp",
    "本地登录态",
    "本地浏览器",
    "本地 chrome",
    "本地Chrome",
    "已有 chrome",
    "已有Chrome",
    "已登录 chrome",
    "已登录Chrome",
    "remote debugging",
    "remote-debugging",
)
HEADED_INTENT_PATTERNS = (
    "有头",
    "可视化",
    "可见浏览器",
    "打开浏览器观察",
    "打开浏览器看",
    "显示浏览器",
    "浏览器窗口",
)
NO_LOGIN_INTENT_PATTERNS = (
    "不需要登录",
    "无需登录",
    "不用登录",
    "免登录",
    "公开页面",
    "公开站点",
    "公开网站",
    "公开url",
    "无需认证",
    "无需鉴权",
    "无需身份",
    "no login",
    "no-login",
    "public page",
    "public site",
)


@dataclass
class QuickRecord:
    case_id: str
    title: str
    query: str
    source: str
    should_run: bool = True
    skip_reason: str = ""
    raw: dict[str, Any] | None = None


def _skill_root() -> Path:
    return Path(__file__).resolve().parents[2]


def _load_env(skill_root: Path) -> Path | None:
    # UIWarden injects per-run browser coordinates into the agent process.
    # The skill .env contains a team-wide external-CDP default (usually 9222)
    # and must never overwrite an embedded run's random WebSocket endpoint.
    protected_runtime = {
        key: value
        for key, value in os.environ.items()
        if key.startswith("UIWARDEN_") or key == "BU_LOCAL_BROWSER_CDP_URL"
    }
    raw = (os.environ.get("DOTENV_PATH") or "").strip()
    candidates = [Path(raw).expanduser()] if raw else []
    candidates.append(skill_root / "config" / ".env")
    for p in candidates:
        if p.is_file():
            from dotenv import dotenv_values, load_dotenv

            load_dotenv(p, override=True)
            os.environ.update(protected_runtime)
            os.environ["DOTENV_PATH"] = str(p)
            _normalize_quick_credential_env(dotenv_values(p))
            return p
    return None


def _log_browser_profile_selection() -> None:
    browser_types = (os.environ.get("BU_BROWSER_TYPES") or "").strip()
    browser_type = (os.environ.get("BU_BROWSER_TYPE") or "chrome").strip()
    selected = browser_types or browser_type
    path_keys = {
        "firefox": "BU_FIREFOX_PATH",
        "custom": "BU_CUSTOM_BROWSER_PATH",
        "edge": "BU_EDGE_PATH",
        "360": "BU_360_BROWSER_PATH",
    }
    path_key = path_keys.get(browser_type.split(":", 1)[0].lower())
    executable = (os.environ.get(path_key) or "").strip() if path_key else ""
    detail = f"，{path_key}={executable}" if path_key and executable else ""
    if browser_type.split(":", 1)[0].lower() == "firefox":
        detail += "（本机路径仅用于检测，执行使用 Playwright 配套 Firefox）"
    _log(f"浏览器 Profile 配置：{selected}{detail}")


def _enforce_uiwarden_browser_execution() -> dict[str, Any] | None:
    session_file = (os.environ.get("UIWARDEN_EMBEDDED_BROWSER_SESSION_FILE") or "").strip()
    kernel = (os.environ.get("UIWARDEN_BROWSER_KERNEL") or "").strip()
    if not session_file and kernel not in {"uiwarden-embedded", "gui-computer"}:
        return None
    if not session_file:
        raise RuntimeError("UIWarden 可控浏览器会话缺少执行会话文件，禁止回退到其它浏览器")
    try:
        config = json.loads(Path(session_file).read_text(encoding="utf-8"))
    except Exception as error:
        raise RuntimeError(f"UIWarden 可控浏览器执行会话文件不可读：{error}") from error
    if config.get("version") != 1 or not config.get("targetId"):
        raise RuntimeError("UIWarden 可控浏览器执行会话协议不兼容或缺少 targetId")
    cdp_url = str(config.get("cdpUrl") or "").strip()
    probe_url = _normalize_uiwarden_cdp_probe_url(cdp_url)
    # browser-use CLI 3.1.6 probes BU_LOCAL_BROWSER_CDP_URL with urllib before
    # it constructs BrowserConfig. The local HTTP base lets Playwright resolve
    # the actual browser WebSocket from /json/version instead of fabricating a
    # WebSocket path from the page target ID.
    os.environ["BU_LOCAL_BROWSER_CDP_URL"] = probe_url
    os.environ["UIWARDEN_EMBEDDED_BROWSER_CDP_URL"] = probe_url
    _log(
        f"UIWarden 可控浏览器协议已锁定：targetId={config['targetId']}，"
        f"cdp={probe_url}"
    )
    return config


def _normalize_uiwarden_cdp_probe_url(value: str) -> str:
    parsed = urlsplit(value)
    if parsed.hostname not in {"127.0.0.1", "localhost"} or not parsed.port:
        raise RuntimeError("UIWarden 可控浏览器 CDP 地址无效，禁止使用 .env 默认 CDP")
    if parsed.scheme in {"ws", "wss"}:
        if not parsed.path.startswith("/devtools/browser/"):
            raise RuntimeError("UIWarden 可控浏览器 CDP WebSocket 地址无效")
    elif parsed.scheme == "http":
        if parsed.path not in {"", "/"}:
            raise RuntimeError("UIWarden 可控浏览器 CDP HTTP 基址无效")
    else:
        raise RuntimeError("UIWarden 可控浏览器 CDP 协议无效")
    return f"http://127.0.0.1:{parsed.port}"


def _normalize_quick_credential_env(values: dict[str, str | None]) -> None:
    """Keep browser-use quick credentials scoped to this skill's .env.

    browser-use also loads its own project .env files. Those files may define
    higher-priority aliases such as QUICK_NAME/QUICK_PASSWORK, so explicitly
    mirror the account configured in this skill into the aliases consumed by
    browser_use.agent.quick before spawning the CLI subprocess.
    """

    username = _first_dotenv_value(
        values,
        "QUICK_NAME",
        "QUICK_USERNAME",
        "BU_USERNAME",
        "USERNAME",
        "user",
    )
    password = _first_dotenv_value(
        values,
        "QUICK_PASSWORK",
        "QUICK_PASSWORD",
        "BU_PASSWORD",
        "PASSWORD",
        "pass",
    )
    if username:
        for key in ("QUICK_NAME", "QUICK_USERNAME", "USERNAME", "user"):
            os.environ[key] = username
        _log(f"登录账号来源已锁定为当前 skill .env：{_mask_credential(username)}")
    if password:
        for key in ("QUICK_PASSWORK", "QUICK_PASSWORD", "PASSWORD", "pass"):
            os.environ[key] = password


def _first_dotenv_value(values: dict[str, str | None], *keys: str) -> str:
    for key in keys:
        value = (values.get(key) or "").strip()
        if value:
            return value
    return ""


def _mask_credential(value: str) -> str:
    if "@" in value:
        name, domain = value.split("@", 1)
        if len(name) <= 2:
            masked = name[:1] + "*"
        else:
            masked = name[:2] + "*" * max(1, len(name) - 4) + name[-2:]
        return f"{masked}@{domain}"
    if len(value) <= 4:
        return "*" * len(value)
    return value[:2] + "*" * (len(value) - 4) + value[-2:]


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("<br>", "\n").replace("<br/>", "\n").strip())


def _log(message: str) -> None:
    print(f"[quick] {message}", flush=True)


def _has_intent(text: str, patterns: Iterable[str]) -> bool:
    blob = (text or "").lower()
    return any(pattern.lower() in blob for pattern in patterns)


def _clip_for_log(text: str, *, limit: int = LOG_QUERY_LIMIT) -> str:
    clean = (text or "").strip()
    clean = re.sub(r"\n{3,}", "\n\n", clean)
    if len(clean) <= limit:
        return clean
    return clean[:limit].rstrip() + f"\n...（已截断，完整内容见产物；原长 {len(clean)} 字符）"


def _log_query_block(*, case_id: str, attempt_no: int, query: str, query_path: Path) -> None:
    body = _clip_for_log(query, limit=LOG_FULL_QUERY_LIMIT)
    _log(
        f"===== quick query BEGIN case={case_id} attempt={attempt_no} =====\n"
        f"{body}\n"
        f"===== quick query END case={case_id} attempt={attempt_no} =====\n"
        f"完整 query 文件：{query_path}"
    )


def _quick_attempt_artifacts(log_dir: Path) -> dict[str, str]:
    result_files = sorted(log_dir.glob("quick_task_result*.json"))
    report_files = sorted(log_dir.glob("quick_task_report*.html"))
    yaml_files = sorted(log_dir.glob("quick_task_*.yaml"))
    return {
        "result_json": str(result_files[-1]) if result_files else "",
        "report_html": str(report_files[-1]) if report_files else "",
        "task_yaml": str(yaml_files[-1]) if yaml_files else "",
    }


def _quick_attempt_result_summary(log_dir: Path, *, returncode: int, dry_run: bool) -> tuple[str, dict[str, str]]:
    artifacts = _quick_attempt_artifacts(log_dir)
    if dry_run:
        return f"dry-run，未实际执行 browser-use quick；returncode={returncode}", artifacts
    result_path = artifacts.get("result_json")
    if not result_path:
        return f"未找到 quick_task_result*.json；returncode={returncode}", artifacts
    try:
        doc = json.loads(Path(result_path).read_text(encoding="utf-8"))
    except Exception as exc:
        return f"读取结果 JSON 失败：{type(exc).__name__}: {exc}；returncode={returncode}", artifacts

    fields: list[str] = [f"returncode={returncode}"]
    for key in ("success", "status", "done", "result", "final_result"):
        if key in doc:
            fields.append(f"{key}={_clean_cell(doc.get(key))}")
    for key in ("message", "error", "reason", "summary"):
        value = _clean_cell(doc.get(key))
        if value:
            fields.append(f"{key}={value[:700]}")
            break
    if len(fields) == 1:
        fields.append("结果文件存在，但未识别到 success/status/message 等字段")
    return "；".join(fields), artifacts


def _normalize_case_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(f"{idx}. {_clean_cell(v)}" for idx, v in enumerate(value, start=1) if _clean_cell(v))
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, indent=2)
    return _clean_cell(value)


def _split_query_sentences(text: str) -> list[str]:
    return [part.strip() for part in re.split(r"(?<=。)", text or "") if part.strip()]


def _same_url(a: str, b: str) -> bool:
    if not a or not b:
        return False
    try:
        pa = urlsplit(a)
        pb = urlsplit(b)
        qa = urlencode(sorted(parse_qsl(pa.query, keep_blank_values=True)))
        qb = urlencode(sorted(parse_qsl(pb.query, keep_blank_values=True)))
        na = urlunsplit((pa.scheme.lower(), pa.netloc.lower(), pa.path.rstrip("/"), qa, ""))
        nb = urlunsplit((pb.scheme.lower(), pb.netloc.lower(), pb.path.rstrip("/"), qb, ""))
        return na == nb
    except Exception:
        return a.rstrip("/").lower() == b.rstrip("/").lower()


def _entry_url_from_query(query: str) -> str:
    m = re.search(r"入口\s*URL[：:]\s*(https?://[^\s。；;，,]+)", query or "", re.I)
    return m.group(1).strip() if m else ""


def _dedupe_entry_hint(query: str) -> str:
    """Remove redundant "建议入口" when it repeats the explicit entry URL.

    Batch records can contain both "入口 URL：..." and "建议入口：访问入口：...".
    That inflates quick prompts and makes the model treat both as independent
    navigation goals. Keep one authoritative entry.
    """
    entry_url = _entry_url_from_query(query)
    if not entry_url:
        return query
    parts: list[str] = []
    for sent in _split_query_sentences(query):
        clean = sent.rstrip("。")
        if clean.startswith("建议入口"):
            hint_url = extract_url(clean)
            if hint_url and _same_url(hint_url, entry_url):
                continue
        parts.append(sent)
    return "".join(parts).strip()


def _needs_text_timestamp_rule(query: str) -> bool:
    if not query or TEXT_TIMESTAMP_RULE_MARKER in query:
        return False
    return any(pattern.lower() in query.lower() for pattern in TEXT_INPUT_HINT_PATTERNS)


def _append_text_timestamp_rule(query: str, text_timestamp: str) -> str:
    """Tell quick to make only agent-invented text values unique.

    This avoids duplicate names for records that say "输入名称" without a
    concrete value, while preserving user-specified strings used in assertions.
    """
    query = (query or "").strip()
    if not text_timestamp or not _needs_text_timestamp_rule(query):
        return query
    rule = (
        f"{TEXT_TIMESTAMP_RULE_MARKER}：凡步骤中需要输入名称、标题、描述、备注、问题、答案、标签名、库名等文本，"
        f"且用例或用户未明确给定具体内容时，请自行生成简短可识别文本，并在文本末尾追加本次时间戳 {text_timestamp}；"
        "若用例或用户已明确给出具体文本、URL、文件路径、账号、密码或断言期望值，必须原样使用，不要追加或改写。"
    )
    if "\n" in query:
        sep = "\n\n"
    elif query.endswith(("。", "；", ";", ".", "！", "!", "？", "?")):
        sep = ""
    else:
        sep = "。"
    return f"{query}{sep}{rule}"


def _has_upload_intent(query: str) -> bool:
    text = query or ""
    markers = ("上传", "选择文件", "本地文件", "upload", "file input")
    return any(marker.lower() in text.lower() for marker in markers)


def _has_explicit_upload_file_path(query: str) -> bool:
    return bool(_extract_explicit_upload_file_paths(query))


def _extract_explicit_upload_file_paths(query: str) -> list[str]:
    if not query:
        return []
    file_suffixes = (
        "pdf",
        "jpg",
        "jpeg",
        "png",
        "gif",
        "webp",
        "md",
        "txt",
        "csv",
        "tsv",
        "xlsx",
        "xls",
        "doc",
        "docx",
        "zip",
    )
    suffix_pattern = "|".join(file_suffixes)
    path_pattern = rf"(?<!\S)(?:/|~/|\.{{1,2}}/)[^\s\"'，,；;）)]+\.({suffix_pattern})(?=$|[\s\"'，,；;）)])"
    paths: list[str] = []
    seen: set[str] = set()
    for match in re.finditer(path_pattern, query, flags=re.IGNORECASE):
        path = match.group(0).strip()
        key = path.lower()
        if key not in seen:
            seen.add(key)
            paths.append(path)
    return paths


def _append_explicit_upload_path_rule(query: str) -> str:
    query = (query or "").strip()
    if not query or UPLOAD_EXPLICIT_PATH_RULE_MARKER in query:
        return query
    paths = _extract_explicit_upload_file_paths(query)
    if not paths or not _has_upload_intent(query):
        return query
    existing = [p for p in paths if Path(p).expanduser().is_file()]
    listed = "；".join(existing or paths)
    existence = "这些路径已在本机存在。" if existing else "若路径执行时不存在，应报告路径问题，不得臆造替代路径。"
    rule = (
        f"{UPLOAD_EXPLICIT_PATH_RULE_MARKER}：本任务已经明确指定上传文件路径：{listed}。"
        "执行 upload 动作时，path 字段必须从上述路径中逐字复制一个；"
        "禁止使用 /path/to/local/image.jpg、/path/to/、example、sample、test.jpg 等示例或占位路径，"
        "也禁止把明确路径改成相对路径。"
        f"{existence}"
    )
    sep = "\n\n" if "\n" in query else ("。 " if not query.endswith(("。", ".", "！", "!", "？", "?")) else " ")
    return f"{query}{sep}{rule}"


def _append_upload_fixture_hint(query: str, *, skill_root: Path) -> str:
    query = (query or "").strip()
    if not query or UPLOAD_FIXTURE_RULE_MARKER in query:
        return query
    if not _has_upload_intent(query) or _has_explicit_upload_file_path(query):
        return query
    upload_dir = skill_root / "config" / UPLOAD_FILES_DIRNAME
    fixtures = [
        ("图片/JPG", upload_dir / "file_jpg.jpg"),
        ("图片/JPEG", upload_dir / "file_jpeg.jpeg"),
        ("PDF/文档", upload_dir / "file_pdf.pdf"),
        ("Markdown/文本", upload_dir / "file_md.md"),
    ]
    existing = [(label, path) for label, path in fixtures if path.is_file()]
    if not existing:
        return query
    choices = "；".join(f"{label}: {path}" for label, path in existing)
    rule = (
        f"{UPLOAD_FIXTURE_RULE_MARKER}：用户未明确指定上传文件路径；若本任务需要上传文件，"
        f"请根据页面文案、accept 类型或业务语义从以下内置样例中选择一个并使用完整路径上传：{choices}。"
        "若页面明确限制的类型与这些样例均不匹配，不要扫描其它目录，记录未覆盖/需人工核对。"
    )
    sep = "\n\n" if "\n" in query else ("。 " if not query.endswith(("。", ".", "！", "!", "？", "?")) else " ")
    return f"{query}{sep}{rule}"


def normalize_record_query_for_execution(
    record: QuickRecord,
    *,
    task_prefix: str = "",
    text_timestamp: str = "",
) -> QuickRecord:
    query = sanitize_quick_query_for_cli(record.query or "")
    query = _dedupe_entry_hint(query)
    if task_prefix:
        prefix_url = extract_url(task_prefix)
        if prefix_url and not _entry_url_from_query(query) and not extract_url(query):
            query = "。".join(p for p in (task_prefix, query) if p)
    if not _is_health_record(record):
        query = _append_text_timestamp_rule(query, text_timestamp)
    query = _append_explicit_upload_path_rule(query)
    return QuickRecord(
        case_id=record.case_id,
        title=record.title,
        query=query,
        source=record.source,
        should_run=record.should_run,
        skip_reason=record.skip_reason,
        raw=record.raw,
    )


def _strip_smoke_markers(text: str) -> str:
    clean = (text or "").strip()
    for marker in SMOKE_WALKTHROUGH_MARKERS:
        clean = clean.replace(marker, "")
    clean = re.sub(r"\bQUICK_URL\s*=\s*https?://\S+", "", clean, flags=re.IGNORECASE)
    clean = re.sub(r"https?://[^\s<>\"'，,。；;）)]+", "", clean, flags=re.IGNORECASE)
    clean = re.sub(r"^(执行|帮我|请|对|针对|给|做|进行)\s*", "", clean)
    clean = re.sub(r"(的|一下|下)\s*$", "", clean)
    clean = re.sub(r"^(的|这个|该)\s*", "", clean)
    clean = re.sub(r"\s+", " ", clean).strip(" ，。；;")
    clean = clean.replace("执行  的", "").strip(" ，。；;")
    return clean


def smoke_focus_from_query(*, user_query: str, target_url: str = "", task_prefix: str = "") -> str:
    raw = (user_query or "").strip()
    url = target_url or extract_url(raw) or resolve_target_url(raw, task_prefix=task_prefix)
    focus = _strip_smoke_markers(raw)
    if url:
        focus = focus.replace(url, "").strip(" ，。；;")
    if task_prefix:
        focus = focus.replace(task_prefix, "").strip(" ，。；;")
    return focus.strip(" ，。；;")


def is_custom_smoke_focus(focus: str) -> bool:
    clean = re.sub(r"https?://\S+", "", focus or "", flags=re.IGNORECASE)
    clean = re.sub(r"\b(cdp|cdp模式|no-cdp|非cdp|headless|headed)\b", "", clean, flags=re.IGNORECASE)
    for token in (
        "有头模式",
        "无头模式",
        "有头",
        "无头",
        "打开浏览器",
        "浏览器窗口",
        "快速",
        "快测",
        "一下",
        "测试",
        "测一下",
    ):
        clean = clean.replace(token, "")
    clean = clean.strip(" ，。；;,.")
    if not clean:
        return False
    custom_markers = (
        "重点",
        "关注",
        "校验",
        "断言",
        "必须",
        "需要覆盖",
        "验证",
        "检查",
        "确认",
        "模块",
        "按钮",
        "表单",
        "输入框",
        "搜索",
        "筛选",
        "红色",
        "告警",
        "报错",
        "接口",
        "性能",
        "登录",
        "下单",
        "创建",
        "新增",
        "提交",
        "支付",
    )
    return any(marker in clean for marker in custom_markers) or len(clean) >= 8


def build_smoke_walkthrough_query(*, user_query: str, target_url: str = "", task_prefix: str = "") -> str:
    """Build an aiQuick-friendly smoke/walkthrough prompt for visible main modules."""
    raw = (user_query or "").strip()
    url = target_url or extract_url(raw) or resolve_target_url(raw, task_prefix=task_prefix)
    focus = smoke_focus_from_query(user_query=raw, target_url=url, task_prefix=task_prefix)
    focus = focus or "可见一级主模块与核心主页面的前端可用性"
    exclusive_focus = any(
        marker in raw
        for marker in (
            "不做其他模块",
            "不要做其他模块",
            "不验证其他模块",
            "不要验证其他模块",
            "不检查其他模块",
            "不要检查其他模块",
            "只验证",
            "仅验证",
            "只看",
            "仅看",
            "只测",
            "仅测",
        )
    )
    parsed = urlsplit(url) if url else None
    origin = f"{parsed.scheme}://{parsed.netloc}" if parsed and parsed.scheme and parsed.netloc else ""

    url_line = f"被测入口：{url}" if url else "被测入口：使用本次 quick 解析到的目标 URL"
    origin_line = f"可回退入口：{origin}" if origin else "可回退入口：同一站点首页或本次入口页"
    domain_line = (
        f"同域边界：本次冒烟主范围限定在 {parsed.netloc} 及其同页面路由内；"
        "每次点击、提交、打开新页后都要先检查地址栏 host；若 host 不是该域名，立即记录外部页标题/是否白屏/是否报错，"
        "然后返回可回退入口继续其它本域模块，不要在外部页继续点击、登录、展开导航或重跑同一路径。"
        if parsed and parsed.netloc
        else "同域边界：本次冒烟主范围限定在入口页所属站点；若跳转到不同域名/不同平台，只记录可达性，然后返回入口继续其它模块，不继续探索外部平台。"
    )
    scope_line = (
        "排他范围硬约束：用户已明确只验证指定模块/流程或不做其他模块验证；本次只允许盘点无关模块名称，不得点击进入、不得做轻量交互、不得把无关模块写成未覆盖。"
        if exclusive_focus
        else "自定义范围规则：若用户附加关注点明确指定某个模块、控件、对话、校验或断言，本次冒烟以该指定范围为主；其它未被用户要求的模块只需在模块清单中记录可见性，不要为了覆盖率强行深挖，也不要把未进入这些无关模块写成未覆盖问题。"
    )
    module_walkthrough_line = (
        "3. 本次为排他范围冒烟：除用户指定模块/流程外，其它可见一级主模块只记录到模块清单，不点击进入、不验证功能、不列为未覆盖；直接进入用户指定模块完成核心检查。"
        if exclusive_focus
        else "3. 按模块清单逐个做轻量冒烟：进入每个可见一级主模块或其代表性子入口，等待页面稳定，确认模块主内容、关键列表/卡片/表单/搜索区/操作按钮是否可见，核心文案是否明显缺失或错位。"
    )
    focus_priority_line = (
        "8. 用户附加关注点为排他范围时，优先且只覆盖这些点；视觉盘点后立即进入指定模块/控件/路径，不要先验证其它可见一级主模块。"
        if exclusive_focus
        else "8. 若用户附加关注点包含具体模块/控件/路径，优先覆盖这些点；但仍需先盘点导航模块，并尽量覆盖其他可见一级主模块。"
    )
    return sanitize_quick_query_for_cli(
        "\n".join(
            [
                "请执行一次前端冒烟验证 / 页面走查，覆盖入口页可见的一级主模块与核心主页面前端可用性，不做破坏性写操作。",
                "重要完成规则：页面稳定、登录成功、看到默认模块都只是前置步骤，不代表任务完成；在完成模块清单盘点、逐模块覆盖、疑似问题/未覆盖汇总之前，禁止调用 done。",
                "如果当前只是观察到页面稳定、模块可见、等待完成或可进入下一步，请继续执行下一步浏览器动作（点击、输入、滚动、等待、回退等），不要用 done 表示阶段完成。",
                url_line,
                origin_line,
                domain_line,
                "登录策略：使用 config/.env 账号密码；若 CDP 已登录可跳过登录。",
                f"用户附加关注点：{focus}",
                scope_line,
                "",
                "走查范围：",
                "1. 打开入口并等待主页面稳定，确认不是白屏、崩溃页、登录死循环、404/500、无权限页或长期 loading。",
                "2. 先做一次视觉模块盘点：只根据当前页面实际可见内容读取左侧导航、顶部导航、可展开分组和当前页面标题，列出所有可见一级主模块；不要套用历史站点、固定产品或参考资料里的模块名，不要虚构不可见模块；不要只停留在默认打开的第一个模块。",
                module_walkthrough_line,
                "4. 每个模块只做一到两个只读或可撤销的核心交互，例如切换 tab、展开筛选/搜索区、打开列表详情后返回、展开配置面板；不提交新增、删除、发布、支付等破坏性动作。",
                "5. 若模块存在搜索框、筛选输入框、试用对话框、评论/问题输入框等低风险输入场景，应输入一段简短测试文本并点击对应查询/发送/提交/回车，验证输入值可见、按钮响应、结果区/提示/校验状态有合理变化；提交前若按钮语义是新增、保存、发布、删除、购买、支付、开通等会产生业务写入或付费动作，则不要提交，只记录未覆盖原因。",
                "6. 导航点击可靠性：若点击某个导航/菜单后页面未变化，先怀疑执行器未点准或命中折叠/遮挡区域；必须换一种方式重试，例如点击文字中心、点击链接文本、滚动后再点、展开父级、用键盘 Enter、或先点相邻项再回来。只有在明确点中了正确控件且页面给出错误/禁用/异常提示时，才可写入“疑似问题”；单次或定位不确定的“点击无响应”只能写入“未覆盖：执行定位未确认/未点准”，不要推断为链接失效或页面 bug。",
                "7. 重点观察每个模块上的疑似 bug：红色错误/警告、异常 toast/modal/banner、接口失败提示、空数据异常、按钮无响应、组件遮挡重叠、内容溢出、关键控件不可点击、明显视觉错乱。",
                focus_priority_line,
                "9. 验证码、滑块、人机校验、安全验证、登录二次校验等风控页面通常是正常防护机制，不要直接判定为前端 bug 或整体阻塞；记录到“未覆盖/风险说明”，然后优先回到可回退入口或入口页，继续走查其他可见主模块。",
                "10. 如果某个模块因为权限、风控、验证码、数据缺失或单点异常无法继续，不要反复尝试绕过；标记该模块未覆盖或疑似问题后继续检查其他可访问主导航/主模块，最后统一汇总总报告。",
                "11. 跨域/外部平台处理：若某个导航项、按钮、表单提交或顶部入口跳转到与被测入口不同域名、不同控制台或明显外部平台，不要把外部平台的导航当成本次模块清单继续展开；只记录“模块/入口可跳转到外部平台 + 外部页面标题/是否白屏/是否报错”，随后立即返回被测入口或上一页继续检查本域其它模块。外部页面上的二级导航、合同/订单/文档等不要继续深挖，也不要反复触发同一跨域入口。",
                "",
                "完成后必须在最终 done 文本中输出如下结构，便于自动汇总：",
                "冒烟模式禁止阶段性 done(success=false)：除非浏览器崩溃或全站完全不可达，否则即使某个目标模型、选项、按钮或模块暂时未找到，也必须继续用搜索、滚动列表容器、键盘方向键/Enter、清空筛选词后重试、关闭下拉重开等方式复核；仍未完成时写入“未覆盖/需人工核对”，并输出完整 SMOKE_REPORT，不要提前结束。",
                "如果存在未覆盖、需人工核对或疑似问题，也要先完成可继续执行的其它检查，最后用完整 SMOKE_REPORT 一次性汇总；不要用 done(success=false) 表达某个局部步骤失败。",
                "疑似问题属于冒烟报告结论，不属于执行器失败；除浏览器崩溃、全站完全不可达等无法产出报告的情况外，最终 done 应使用 success=true 并在 SMOKE_REPORT 中写清问题。",
                "异步/生成类任务专项：提交后若出现“生成中、处理中、加载中、AI 正在回复、任务执行中”等状态，不要等待 3-5 秒就定性；应分段等待并复查，累计至少 30 秒或直到出现结果/明确错误提示，再写入报告。",
                "下拉/选择类控件专项：目标选项首屏不可见不等于不存在；必须至少尝试滚动下拉列表到底部/顶部、若有搜索框则输入目标文本搜索、用键盘在下拉内移动、关闭重开后重试。只有完成这些动作后仍不可见，才能记录“未覆盖：目标选项未找到/需人工核对”。",
                "只有当 done 文本以 SMOKE_REPORT 开头并逐行包含下列字段名时，才允许结束任务；字段名必须原样输出，不要用“覆盖情况”等同义词替代“模块覆盖”；不要输出阶段性 done，不要把“可进入下一步”或单个未覆盖项当作完成。",
                "SMOKE_REPORT",
                "结论：通过 / 有疑似问题 / 部分未覆盖",
                "模块清单：列出视觉盘点识别到的所有可见一级主模块；若只识别到当前模块，说明原因",
                "模块覆盖：逐模块列出已进入/已检查/未覆盖及原因，必须体现是否覆盖了默认模块之外的其他模块；跨域入口写“仅验证跳转可达，未深入外部平台”",
                "已检查：用分号列出实际检查过的页面区域、核心交互、输入/提交验证结果",
                "疑似问题：如无问题写“无”；如有问题逐条写【严重级别：高/中/低】【位置】【现象】【复现动作】【证据文本】。验证码/安全验证本身不要写入疑似问题，除非它异常循环、无法关闭且阻止所有主页面访问。",
                "未覆盖：列出因验证码/风控、权限、数据、时间、页面不可达、跨域外部平台不深入等原因未覆盖的模块；没有则写“无”",
                "建议：给出下一步人工复核或专项测试建议；没有则写“无”",
            ]
        )
    )


def _strip_health_markers(text: str) -> str:
    clean = text or ""
    for marker in HEALTH_CHECK_MARKERS:
        clean = clean.replace(marker, "")
    clean = re.sub(r"^(执行|检查|验证|测试|看一下|看下|帮我|对)\s*", "", clean)
    clean = re.sub(r"^(一下|一下子|下)\s*", "", clean)
    clean = re.sub(r"(是否)?(正常打开|正常渲染|正常加载|可访问|能访问)\s*", "", clean)
    clean = re.sub(r"\s+", " ", clean).strip(" ，。；;")
    return clean


def build_health_check_query(*, user_query: str, target_url: str = "", task_prefix: str = "") -> str:
    """Build an aiQuick-friendly page health/liveness prompt."""
    raw = (user_query or "").strip()
    url = target_url or extract_url(raw) or resolve_target_url(raw, task_prefix=task_prefix)
    focus = _strip_health_markers(raw)
    if url:
        focus = focus.replace(url, "").strip(" ，。；;")
    if task_prefix:
        focus = focus.replace(task_prefix, "").strip(" ，。；;")
    focus = focus or "页面是否正常打开、渲染和加载"
    parsed = urlsplit(url) if url else None

    url_line = f"被测入口：{url}" if url else "被测入口：使用本次 quick 解析到的目标 URL"
    domain_line = (
        f"同域边界：本次巡检只关注 {parsed.netloc} 入口页本身；若入口自动跳转到登录页、验证码页或外部域，"
        "只记录跳转后的 URL/标题/是否白屏/是否报错，不继续登录、点击导航或探索其它页面。"
        if parsed and parsed.netloc
        else "同域边界：本次巡检只关注入口页本身；若自动跳转到登录页、验证码页或外部域，只记录状态，不继续探索。"
    )
    return sanitize_quick_query_for_cli(
        "\n".join(
            [
                "请执行一次页面巡检 / 页面探活，只判断入口页面是否可访问、正常渲染和完成基本加载。",
                "重要范围：这不是冒烟走查，不做模块盘点，不探索导航，不验证业务功能，不做输入、提交、新增、删除、购买、发布等操作；页面稳定后即可总结。",
                url_line,
                domain_line,
                f"用户附加关注点：{focus}",
                "",
                "巡检步骤：",
                "1. 打开入口 URL，等待页面稳定到不再明显 loading；最多只做必要等待或刷新一次，不做功能点击。",
                "2. 观察首屏是否有可识别标题、导航、主体内容、卡片/表格/列表/表单或登录表单之一；登录页、验证码页、安全校验页若能正常渲染，视为页面可达但业务内容未覆盖，不判定为前端 bug。",
                "3. 重点检查页面探活异常：DNS/证书/网络错误、404/500/无权限错误页、白屏、崩溃页、长时间 loading、资源加载失败导致主体缺失、明显红色错误/异常 toast/modal/banner、严重布局错乱。",
                "4. 若自动跳到非同域页面，只记录外部 URL、标题、是否白屏/报错，然后结束；不要继续探索外部页面。",
                "",
                "完成后必须在最终 done 文本中输出如下结构，便于自动汇总：",
                "只有当 done 文本以 HEALTH_REPORT 开头并包含下列所有字段时，才允许结束任务。",
                "HEALTH_REPORT",
                "结论：通过 / 异常 / 部分可达",
                "入口：实际访问的入口 URL",
                "最终页面：最终 URL、页面标题或可见主标题",
                "渲染状态：说明是否白屏、是否长期 loading、主体内容是否可见",
                "异常：如无问题写“无”；如有问题逐条写【严重级别：高/中/低】【现象】【证据文本】",
                "备注：登录页/验证码/权限/跨域跳转等非功能深入原因；没有则写“无”",
            ]
        )
    )


def _extract_text_values(value: Any, *, limit: int = 30000) -> list[str]:
    texts: list[str] = []

    def walk(obj: Any) -> None:
        if sum(len(t) for t in texts) > limit:
            return
        if isinstance(obj, str):
            s = _clean_cell(obj)
            if s:
                texts.append(s)
        elif isinstance(obj, dict):
            for v in obj.values():
                walk(v)
        elif isinstance(obj, list):
            for v in obj:
                walk(v)

    walk(value)
    return texts


def _latest_result_doc(attempt: dict[str, Any]) -> dict[str, Any]:
    path = attempt.get("result_json")
    if not path:
        return {}
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except Exception:
        return {}


def _extract_done_texts(doc: dict[str, Any]) -> list[str]:
    texts: list[str] = []
    history = doc.get("history") or []
    if not isinstance(history, list):
        return texts
    for item in history:
        if not isinstance(item, dict):
            continue
        for action in item.get("executed_actions") or []:
            if not isinstance(action, dict) or "done" not in action:
                continue
            payload = action.get("done") or {}
            text = _clean_cell(payload.get("text"))
            if text:
                texts.append(text)
    return texts


def _extract_smoke_section(text: str) -> str:
    marker = "SMOKE_REPORT"
    lines = text.splitlines()
    for idx, line in enumerate(lines):
        if line.strip().startswith(marker):
            return "\n".join(lines[idx:]).strip()
    return text.strip()


def _extract_health_section(text: str) -> str:
    marker = "HEALTH_REPORT"
    lines = text.splitlines()
    for idx, line in enumerate(lines):
        if line.strip().startswith(marker):
            return "\n".join(lines[idx:]).strip()
    return text.strip()


def _is_smoke_record(record: QuickRecord) -> bool:
    raw = record.raw or {}
    return raw.get("case_type") == "smoke_walkthrough" or record.case_id.endswith("smoke_walkthrough")


def _is_health_record(record: QuickRecord) -> bool:
    raw = record.raw or {}
    return raw.get("case_type") == "health_check" or record.case_id.endswith("health_check")


def _is_negated_issue_line(line: str) -> bool:
    compact = re.sub(r"\s+", "", line)
    negation_markers = (
        "无红色",
        "无错误",
        "无报错",
        "无异常",
        "无警告",
        "无告警",
        "无白屏",
        "无崩溃",
        "无响应异常",
        "未发现",
        "没有发现",
        "不存在",
        "正常",
        "通过",
    )
    return any(marker in compact for marker in negation_markers)


def _is_normal_smoke_risk_line(line: str) -> bool:
    compact = re.sub(r"\s+", "", line)
    if not any(keyword in compact for keyword in SMOKE_NORMAL_RISK_KEYWORDS):
        return False
    if any(marker in compact for marker in ("未触发", "未出现", "没有触发", "无验证码", "无安全校验", "无安全验证", "无需验证", "无需校验")):
        return False
    if any(marker in compact for marker in ("异常循环", "无法关闭且阻止所有", "阻止所有主页面访问", "全站不可访问")):
        return False
    return True


def _extract_smoke_issues(text: str) -> list[str]:
    issues: list[str] = []
    for raw_line in text.splitlines():
        line = _clean_cell(raw_line.strip("-* 　"))
        if not line:
            continue
        if _is_negated_issue_line(line):
            continue
        if _is_normal_smoke_risk_line(line):
            continue
        if "未覆盖" in line and not any(keyword in line for keyword in ("失败", "错误", "报错", "异常", "400", "500", "failed", "error")):
            continue
        if any(keyword.lower() in line.lower() for keyword in SMOKE_ISSUE_KEYWORDS):
            if "无" in line and len(line) <= 12:
                continue
            issues.append(line[:500])
    # 去重保序
    seen: set[str] = set()
    out: list[str] = []
    for issue in issues:
        key = issue.lower()
        if key not in seen:
            seen.add(key)
            out.append(issue)
    return out


def _extract_smoke_uncovered(text: str) -> list[str]:
    uncovered: list[str] = []
    in_uncovered = False
    for raw_line in text.splitlines():
        line = _clean_cell(raw_line.strip("-* 　"))
        if not line:
            continue
        if line.startswith("未覆盖"):
            in_uncovered = True
            rest = line.split("：", 1)[1] if "：" in line else line.split(":", 1)[1] if ":" in line else ""
            if rest and rest.strip(" 。；;") != "无":
                uncovered.append(rest[:500])
            continue
        if line.startswith(("建议", "疑似问题", "已检查", "结论", "SMOKE_REPORT")):
            in_uncovered = False
        if in_uncovered and line != "无":
            uncovered.append(line[:500])
        elif _is_normal_smoke_risk_line(line):
            uncovered.append(line[:500])
        elif "未覆盖" in line and "无未覆盖" not in re.sub(r"\s+", "", line) and not re.search(r"未覆盖(?:原因)?[：:]?\s*无(?:[。；;]|$)", line):
            uncovered.append(line[:500])
    seen: set[str] = set()
    out: list[str] = []
    for item in uncovered:
        key = item.lower()
        if key not in seen:
            seen.add(key)
            out.append(item)
    return out


def _missing_smoke_required_sections(text: str) -> list[str]:
    compact = re.sub(r"\s+", "", text or "")
    required_aliases = {
        "SMOKE_REPORT": ("SMOKE_REPORT",),
        "模块清单": ("模块清单", "一级主模块", "可见主模块", "可见一级主模块", "盘点当前页可见一级主模块"),
        "模块覆盖": ("模块覆盖", "覆盖情况"),
        "已检查": ("已检查", "检查情况", "实际检查"),
        "疑似问题": ("疑似问题",),
        "未覆盖": ("未覆盖",),
    }
    return [section for section, aliases in required_aliases.items() if not any(alias in compact for alias in aliases)]


def _missing_health_required_sections(text: str) -> list[str]:
    compact = re.sub(r"\s+", "", text or "")
    required = ("HEALTH_REPORT", "结论", "入口", "最终页面", "渲染状态", "异常", "备注")
    return [section for section in required if section not in compact]


def _unique_preserve_order(items: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    invalid_prefixes = (
        "列出",
        "若只",
        "说明",
        "检查",
        "记录",
        "标记",
        "点击",
        "左侧导航",
        "位于",
        "页面",
        "当前",
    )
    for raw in items:
        item = _clean_cell(str(raw or "").strip(" -;；,，。"))
        if not item or item.startswith(invalid_prefixes) or "‘" in item or "’" in item:
            continue
        key = item.lower()
        if key not in seen:
            seen.add(key)
            out.append(item)
    return out


def _read_optional_text(path: str | Path, *, limit: int = 240_000) -> str:
    try:
        p = Path(path)
        if not p.exists() or not p.is_file():
            return ""
        text = p.read_text(encoding="utf-8", errors="ignore")
        if len(text) > limit:
            return text[-limit:]
        return text
    except Exception:
        return ""


def _extract_json_object(text: str) -> dict[str, Any]:
    raw = (text or "").strip()
    raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.I)
    raw = re.sub(r"\s*```$", "", raw)
    start = raw.find("{")
    end = raw.rfind("}")
    if start < 0 or end < start:
        raise ValueError("LLM 输出中未找到 JSON 对象")
    data = json.loads(raw[start : end + 1])
    if not isinstance(data, dict):
        raise ValueError("LLM 输出 JSON 不是对象")
    return data


def _trim_text(text: str, limit: int) -> str:
    text = text or ""
    if len(text) <= limit:
        return text
    head_len = max(0, limit // 3)
    return text[:head_len] + "\n...\n[中间内容已截断]\n...\n" + text[-(limit - head_len) :]


def _attempt_log_summary_from_result(result: dict[str, Any]) -> str:
    """Return a compact log summary for LLM report review.

    Raw browser-use logs contain repeated prompts, planner context and transient
    states. Feeding them directly to the report LLM makes middle states look
    more important than the final done report. Keep final done text and a small
    set of actionable error/progress lines only.
    """
    chunks: list[str] = []
    for attempt in result.get("attempts") or []:
        attempt_no = attempt.get("attempt")
        validation = attempt.get("validation") or {}
        done_text = _clean_cell(validation.get("agent_done_text"))
        if done_text:
            chunks.append(f"[attempt {attempt_no}] final_done:\n{_trim_text(done_text, 6000)}")
        result_json = attempt.get("result_json") or ""
        if not result_json:
            continue
        log_file = Path(result_json).parent / "log.log"
        log_text = _read_optional_text(log_file, limit=220_000)
        if not log_text:
            continue
        selected_lines: list[str] = []
        markers = (
            "SMOKE_REPORT",
            "HEALTH_REPORT",
            "JSONDecodeError",
            "Traceback",
            "Exception",
            "ERROR",
            "Error",
            "错误",
            "报错",
            "异常",
            "失败",
            "超时",
            "timeout",
            "AI 正在回复",
            "AI响应",
            "AI 响应",
        )
        for raw_line in log_text.splitlines():
            line = _clean_cell(raw_line)
            if line and any(marker in line for marker in markers):
                selected_lines.append(line)
        if selected_lines:
            chunks.append(f"[attempt {attempt_no}] selected_log_lines:\n" + "\n".join(selected_lines[-80:]))
    return "\n".join(chunk for chunk in chunks if chunk)


def _result_json_summary_from_results(results: list[dict[str, Any]]) -> str:
    chunks: list[str] = []
    for result in results:
        chunks.append(
            f"case={result.get('case_id', '')} status={result.get('status', '')} "
            f"attempts={len(result.get('attempts') or [])}"
        )
        for attempt in result.get("attempts") or []:
            validation = attempt.get("validation") or {}
            result_json = attempt.get("result_json") or ""
            doc: dict[str, Any] = {}
            if result_json:
                try:
                    doc = json.loads(Path(result_json).read_text(encoding="utf-8"))
                except Exception as e:
                    chunks.append(f"  attempt={attempt.get('attempt')} result_json_parse_error={type(e).__name__}: {e}")
                    doc = {}
            done_texts = _extract_done_texts(doc) if doc else []
            final_done = done_texts[-1] if done_texts else _clean_cell(validation.get("agent_done_text"))
            history = doc.get("history") if isinstance(doc, dict) else None
            final_url = ""
            result_errors: list[str] = []
            result_extracts: list[str] = []
            if isinstance(history, list):
                for item in history[-12:]:
                    if not isinstance(item, dict):
                        continue
                    state = item.get("state") if isinstance(item.get("state"), dict) else {}
                    final_url = _clean_cell(state.get("url")) or final_url
                    for entry in item.get("result") or []:
                        if not isinstance(entry, dict):
                            continue
                        err = _clean_cell(entry.get("error"))
                        if err:
                            result_errors.append(err)
                        extracted = _clean_cell(entry.get("extracted_content"))
                        if extracted:
                            result_extracts.append(extracted)
            chunks.append(
                "\n".join(
                    [
                        f"  attempt={attempt.get('attempt')} returncode={attempt.get('returncode')} resume={attempt.get('resume')}",
                        f"  validation_passed={validation.get('passed')} validation_method={validation.get('method')} agent_done_success={validation.get('agent_done_success')}",
                        f"  validation_reason={_clean_cell(validation.get('reason'))}",
                        f"  history_len={len(history) if isinstance(history, list) else 0} final_url={final_url}",
                        f"  final_done={_trim_text(final_done, 5000)}",
                        "  recent_result_errors=" + json.dumps(_unique_preserve_order(result_errors)[-8:], ensure_ascii=False),
                        "  recent_result_extracts=" + json.dumps(_unique_preserve_order(result_extracts)[-8:], ensure_ascii=False),
                    ]
                )
            )
    return "\n".join(chunk for chunk in chunks if chunk)


def _final_smoke_success_signal(report_items: list[dict[str, Any]], results: list[dict[str, Any]]) -> bool:
    has_success_result = any(result.get("status") == "success" for result in results)
    if not has_success_result:
        return False
    final_sections = "\n".join(str(item.get("smoke_report_text") or "") for item in report_items)
    compact = re.sub(r"\s+", "", final_sections)
    if not compact or "SMOKE_REPORT" not in compact:
        return False
    positive_markers = ("结论：通过", "结论:通过", "全部检查通过", "无任何疑似问题", "疑似问题：无", "疑似问题:无")
    negative_markers = ("结论：有疑似问题", "结论:有疑似问题", "结论：部分未覆盖", "结论:部分未覆盖")
    return any(marker in compact for marker in positive_markers) and not any(marker in compact for marker in negative_markers)


def _is_tool_or_transient_problem(item: str) -> bool:
    compact = re.sub(r"\s+", "", str(item or "")).lower()
    if not compact or _is_empty_problem_point(compact):
        return True
    transient_markers = (
        "jsondecodeerror",
        "json解析",
        "planner",
        "执行器",
        "工具",
        "browser-use",
        "未在30秒",
        "等待30s",
        "等待30秒",
        "ai正在回复",
        "响应等待",
        "中断流程",
    )
    return any(marker in compact for marker in transient_markers)


def _save_issue_screenshots(results: list[dict[str, Any]], batch_dir: Path, *, has_final_issues: bool = True) -> list[dict[str, str]]:
    if not has_final_issues:
        return []
    issue_markers = ("400", "401", "403", "404", "500", "502", "503", "错误", "报错", "异常", "失败", "failed", "error")
    saved: list[dict[str, str]] = []
    assets_dir = batch_dir / "assets"
    for result in results:
        for attempt in result.get("attempts") or []:
            result_json = attempt.get("result_json") or ""
            if not result_json:
                continue
            log_text = _read_optional_text(Path(result_json).parent / "log.log", limit=240_000)
            try:
                doc = json.loads(Path(result_json).read_text(encoding="utf-8"))
            except Exception:
                continue
            history = doc.get("history") if isinstance(doc, dict) else None
            if not isinstance(history, list):
                continue
            matched_items: list[tuple[int, dict[str, Any], str]] = []
            for idx, item in enumerate(history):
                text = json.dumps(
                    {k: v for k, v in (item or {}).items() if k not in {"state", "screenshot_before_actions"}},
                    ensure_ascii=False,
                    default=str,
                )
                item_has_issue = any(marker.lower() in text.lower() for marker in issue_markers)
                if not item_has_issue:
                    continue
                state = item.get("state") if isinstance(item, dict) else {}
                if any((r.get("is_done") for r in (item.get("result") or []) if isinstance(r, dict))):
                    continue
                screenshot = state.get("screenshot") if isinstance(state, dict) else ""
                if not screenshot:
                    screenshot = item.get("screenshot_before_actions") if isinstance(item, dict) else ""
                if not screenshot:
                    continue
                matched_items.append((idx, item, screenshot))
            if matched_items:
                idx, item, screenshot = matched_items[-1]
                state = item.get("state") if isinstance(item, dict) else {}
                try:
                    assets_dir.mkdir(parents=True, exist_ok=True)
                    name = f"issue_screenshot_{len(saved) + 1:02d}_history_{idx}.png"
                    out = assets_dir / name
                    out.write_bytes(base64.b64decode(screenshot))
                    saved.append(
                        {
                            "path": str(out),
                            "relative_path": f"assets/{name}",
                            "history_index": str(idx),
                            "url": str((state or {}).get("url") or ""),
                            "reason": "命中错误/异常相关执行记录，截取最近的问题页面",
                        }
                    )
                except Exception:
                    continue
    return saved


def _extract_smoke_modules_from_text(text: str) -> tuple[list[str], list[str], list[str]]:
    modules: list[str] = []
    covered: list[str] = []
    unconfirmed: list[str] = []
    module_stop = re.compile(r"[，,、；;]+")
    for match in re.finditer(r"(?:左侧导航栏?可见一级主模块|左侧导航栏可见模块|模块清单(?:包括)?|识别到左侧模块)[：:]\s*([^。\n]+)", text):
        segment = re.sub(r"[【】\[\]]", "", match.group(1))
        for part in re.split(module_stop, segment):
            part = _clean_cell(part.strip("。:："))
            if part and len(part) <= 30 and not part.startswith(("顶部", "当前", "以及")):
                modules.append(part)
    quoted_module = r"[‘'“\"]([^’'”\"]{1,30})[’'”\"]模块"
    modules.extend(re.findall(quoted_module, text))
    state_lines = [
        line
        for line in text.splitlines()
        if "State analysis" in line or "Progress evaluation" in line or "状态分析" in line
    ]
    for line in state_lines:
        if "未成功" in line or "未触发" in line or "未发生" in line:
            continue
        for pattern in (
            r"(?:已成功|成功|当前页面为|页面为|已进入|导航至|跳转至)[^。\n]{0,30}" + quoted_module,
            quoted_module + r"[^。\n]{0,45}(?:稳定加载|内容稳定|主内容.*可见|已完成.*验证|进入与内容可见性验证)",
        ):
            for match in re.finditer(pattern, line):
                groups = [g for g in match.groups() if g]
                if groups:
                    covered.append(groups[-1])
    for pattern in (
        quoted_module + r"[^。\n]{0,80}(?:未覆盖|未成功|未触发|未跳转|无响应|未确认)",
        r"(?:未覆盖|未成功进入|未跳转至|无法进入)[^。\n]{0,30}" + quoted_module,
    ):
        for match in re.finditer(pattern, text):
            groups = [g for g in match.groups() if g]
            if groups:
                unconfirmed.append(groups[-1])
    return (
        _unique_preserve_order(modules),
        _unique_preserve_order(covered),
        _unique_preserve_order(unconfirmed),
    )


def _focus_match_terms(requested_focus: str) -> list[str]:
    """Extract generic, request-local terms used to keep custom smoke analysis scoped."""
    terms: list[str] = []
    stopwords = {
        "的",
        "和",
        "与",
        "并",
        "以及",
        "一下",
        "是否",
        "正常",
        "看",
        "看看",
        "关注",
        "重点",
        "验证",
        "校验",
        "检查",
        "测试",
        "冒烟",
        "走查",
        "页面",
        "模块",
        "功能",
        "流程",
        "入口",
        "控件",
        "按钮",
    }
    for chunk in re.split(r"[，,、；;\s]+", requested_focus or ""):
        chunk = _clean_cell(chunk.strip(" -:：。,.，；;"))
        if not chunk:
            continue
        terms.append(chunk)
        normalized = chunk
        for word in stopwords:
            normalized = normalized.replace(word, "")
        normalized = _clean_cell(normalized.strip(" -:：。,.，；;"))
        if normalized and normalized != chunk:
            terms.append(normalized)
        for part in re.split(r"(?:是否|能否|并且|并|和|与|以及|看|验证|校验|检查|测试)", chunk):
            part = _clean_cell(part.strip(" -:：。,.，；;"))
            if part and part not in stopwords:
                terms.append(part)
    return [term for term in _unique_preserve_order(terms) if len(term) >= 2]


def _is_empty_problem_point(item: str) -> bool:
    compact = re.sub(r"\s+", "", str(item or ""))
    return compact in {"无", "无。", "暂无", "不适用", "未发现明确前端bug", "未发现明确前端问题"}


def _has_explicit_problem_evidence(items: Iterable[str]) -> bool:
    text = "\n".join(str(item or "") for item in items)
    if not text.strip():
        return False
    if all(_is_empty_problem_point(item) for item in items):
        return False
    problem_markers = (
        "错误",
        "报错",
        "异常",
        "失败",
        "不可用",
        "无响应",
        "白屏",
        "崩溃",
        "遮挡",
        "错位",
        "溢出",
        "接口",
        "error",
        "failed",
        "exception",
        "timeout",
        "400",
        "401",
        "403",
        "404",
        "500",
        "502",
        "503",
    )
    lowered = text.lower()
    return any(marker.lower() in lowered for marker in problem_markers)


def build_smoke_ai_analysis(
    *,
    report_items: list[dict[str, Any]],
    results: list[dict[str, Any]],
    issues: list[str],
    uncovered: list[str],
    custom_smoke: bool = False,
    issue_screenshots: list[dict[str, str]] | None = None,
    requested_focus: str = "",
) -> dict[str, Any]:
    failed_count = sum(1 for result in results if result.get("status") == "failed")
    validation_failed = any(
        not ((attempt.get("validation") or {}).get("passed"))
        for result in results
        for attempt in (result.get("attempts") or [])[-1:]
    )
    text_blobs = []
    for item in report_items:
        text_blobs.append(str(item.get("smoke_report_text") or ""))
    text_blobs.extend(_attempt_log_summary_from_result(result) for result in results)
    all_text = "\n".join(text_blobs)
    modules, covered_modules, unconfirmed_modules = _extract_smoke_modules_from_text(all_text)
    if not unconfirmed_modules:
        unconfirmed_modules = _unique_preserve_order(
            re.findall(r"([A-Za-z0-9_\-/\u4e00-\u9fff ]{1,30})模块[^。\n]{0,80}未覆盖", "\n".join(uncovered))
        )
    llm_analysis = build_smoke_ai_analysis_with_llm(
        report_items=report_items,
        results=results,
        issues=issues,
        uncovered=uncovered,
        custom_smoke=custom_smoke,
        requested_focus=requested_focus,
    )
    if llm_analysis:
        final_success_signal = _final_smoke_success_signal(report_items, results)
        if custom_smoke and requested_focus:
            focus_terms = _focus_match_terms(requested_focus)
            def related_to_focus(item: str) -> bool:
                return any(term and term in item for term in focus_terms)
            empty_markers = {"无", "无。", "暂无", "不适用"}
            removed = [
                str(item)
                for item in (llm_analysis.get("manual_review_points") or [])
                if str(item).strip() not in empty_markers and not related_to_focus(str(item))
            ]
            if removed:
                kept = [str(item) for item in (llm_analysis.get("manual_review_points") or []) if related_to_focus(str(item))]
                llm_analysis["manual_review_points"] = kept or ["无"]
                ignored = list(llm_analysis.get("ignored_unrequested_scope") or [])
                ignored.append("未被用户指定的其它模块未作为本次冒烟结论依据：" + "；".join(removed))
                llm_analysis["ignored_unrequested_scope"] = _unique_preserve_order(ignored)
        problem_points = [str(item) for item in (llm_analysis.get("problem_points") or [])]
        has_problem_evidence = bool(issues) or _has_explicit_problem_evidence(problem_points)
        if (
            final_success_signal
            and llm_analysis.get("verdict") != "冒烟通过"
            and not issues
            and all(_is_tool_or_transient_problem(item) for item in problem_points)
        ):
            manual_points = [
                str(item)
                for item in (llm_analysis.get("manual_review_points") or [])
                if str(item).strip() and not _is_empty_problem_point(str(item))
            ]
            if problem_points and not all(_is_empty_problem_point(item) for item in problem_points):
                manual_points.append("执行过程中出现过工具/异步中间态噪声，但最终 SMOKE_REPORT 已完成且通过；相关信息仅作为人工抽查参考。")
            llm_analysis["verdict"] = "冒烟通过"
            llm_analysis["reason"] = (
                "最终 SMOKE_REPORT 已通过，且未提取到最终报告中的明确前端 bug；"
                "中间执行器 JSON 解析、短时等待或异步回复状态不再覆盖最终成功结论。"
            )
            llm_analysis["problem_points"] = ["未发现明确前端 bug"]
            llm_analysis["manual_review_points"] = _unique_preserve_order(manual_points) or ["无"]
            has_problem_evidence = False
        if llm_analysis.get("verdict") == "有问题" and not has_problem_evidence:
            manual_points = [
                str(item)
                for item in (llm_analysis.get("manual_review_points") or [])
                if str(item).strip() and not _is_empty_problem_point(str(item))
            ]
            if uncovered:
                manual_points.extend(str(item) for item in uncovered)
            if failed_count or validation_failed:
                manual_points.append("执行器早停、报告字段缺失或完成度验收未通过，需以 HTML/日志回放或 CDP 复核实际页面。")
            llm_analysis["verdict"] = "需要人工核对"
            llm_analysis["reason"] = (
                "未发现明确前端 bug 证据；当前主要风险来自执行器早停、目标项未充分探索或报告验收失败，"
                "不能直接判定页面有问题，需要人工或 CDP 复核。"
            )
            llm_analysis["problem_points"] = ["未发现明确前端 bug"]
            llm_analysis["manual_review_points"] = _unique_preserve_order(manual_points) or ["无"]
        llm_analysis.setdefault("execution_failed_count", failed_count)
        llm_analysis.setdefault("validation_failed", validation_failed)
        llm_analysis.setdefault("custom_smoke", custom_smoke)
        llm_analysis.setdefault("issue_screenshots", issue_screenshots or [])
        return llm_analysis

    if issues:
        verdict = "有问题"
        reason = "执行中发现疑似前端问题，需要进入缺陷复核。"
    elif custom_smoke and (uncovered or failed_count or validation_failed):
        verdict = "需要人工核对"
        reason = "这是带明确关注点/校验要求的自定义冒烟；虽未发现明确前端 bug，但存在未覆盖项、执行定位未确认或报告验收失败，不能判定冒烟通过。"
    else:
        verdict = "冒烟通过"
        if uncovered or failed_count or validation_failed:
            reason = "默认冒烟未发现明确前端 bug；未覆盖项、定位未确认或报告验收失败已列为人工核对项，不影响默认冒烟通过结论。"
        else:
            reason = "关键页面已覆盖且未发现疑似问题或未覆盖项。"
    coverage_points: list[str] = []
    if modules:
        coverage_points.append("识别到模块：" + "、".join(modules))
    if covered_modules:
        coverage_points.append("已实际进入/检查：" + "、".join(covered_modules))
    if not coverage_points:
        coverage_points.append("未能从日志稳定提取覆盖点，请查看 HTML 报告。")
    problem_points = issues[:] if issues else ["未发现明确前端 bug"]
    manual_points: list[str] = []
    if uncovered:
        for item in uncovered:
            if "SMOKE_REPORT" in item and ("未检查" in item or "总体进度" in item):
                manual_points.append("原始 SMOKE_REPORT 与执行日志存在不一致，需以 HTML/日志回放核对实际覆盖过程。")
            else:
                manual_points.append(item)
    if unconfirmed_modules:
        manual_points.append("需人工核对模块/交互：" + "、".join(unconfirmed_modules))
    if failed_count or validation_failed:
        manual_points.append("执行器或报告验收未通过，需以 HTML/日志回放核对实际覆盖过程。")
    manual_points = _unique_preserve_order(manual_points) or ["无"]
    return {
        "verdict": verdict,
        "reason": reason,
        "modules": modules,
        "covered_modules": covered_modules,
        "unconfirmed_modules": unconfirmed_modules,
        "coverage_points": coverage_points,
        "problem_points": problem_points,
        "manual_review_points": manual_points,
        "execution_failed_count": failed_count,
        "validation_failed": validation_failed,
        "custom_smoke": custom_smoke,
        "analysis_method": "rules",
        "issue_screenshots": issue_screenshots or [],
    }


async def _build_smoke_ai_analysis_with_llm_async(
    *,
    report_items: list[dict[str, Any]],
    results: list[dict[str, Any]],
    issues: list[str],
    uncovered: list[str],
    custom_smoke: bool,
    requested_focus: str,
) -> dict[str, Any]:
    try:
        from langchain_core.messages import HumanMessage, SystemMessage
    except ImportError as e:
        raise RuntimeError("LLM 冒烟分析需要 langchain_core") from e
    llm = _get_doc_llm()
    original_cases = "\n\n".join(
        _read_optional_text(result.get("original_case_path") or "", limit=12_000)
        for result in results
        if result.get("original_case_path")
    )
    done_reports = "\n\n".join(str(item.get("smoke_report_text") or "") for item in report_items)
    logs = "\n\n".join(_attempt_log_summary_from_result(result) for result in results)
    result_json_summary = _result_json_summary_from_results(results)
    system = (
        "你是资深前端测试报告分析专家。请基于执行日志整体理解冒烟结果，不要只靠关键字。\n"
        "只输出 JSON 对象，不要 markdown。字段固定：verdict, reason, coverage_points, problem_points, manual_review_points, ignored_unrequested_scope, confidence。\n"
        "verdict 只能是：冒烟通过、需要人工核对、有问题。\n"
        "判断规则：若用户明确要求的模块/交互/校验点出现错误、失败响应、未完成或无法验证，verdict=有问题或需要人工核对；"
        "只有存在明确页面错误、接口错误、错误码、失败提示、白屏崩溃、可复现交互异常等产品侧证据时，verdict=有问题。"
        "如果只是执行器早停、报告字段缺失、未滚动/未搜索导致目标项未确认、定位未确认或探索不充分，verdict=需要人工核对，不要判成有问题。"
        "若最终 Agent SMOKE_REPORT 已明确通过、最终 done success=true 且没有最终报告中的疑似问题，"
        "中间日志中的等待、AI 正在回复、短时超时、planner JSONDecodeError、工具中断等只能作为人工抽查信息，不能覆盖最终成功结论。"
        "默认冒烟没有明确 bug 时可为冒烟通过。\n"
        "范围规则：用户明确只要求某模块或某交互时，不要把未要求的其它模块写成问题或人工核对项；必须放入 ignored_unrequested_scope。"
        "如果原始用例里含脚本自动生成的通用模块覆盖步骤，但用户附加关注点更具体，必须以用户附加关注点作为本次必须验收范围。"
    )
    user = (
        f"是否自定义冒烟：{custom_smoke}\n"
        f"用户明确关注点/必须验收范围：{requested_focus or '(默认冒烟，无额外指定)'}\n\n"
        f"原始用例/用户要求：\n{_trim_text(original_cases, 16000)}\n\n"
        f"Agent 最终 SMOKE_REPORT：\n{_trim_text(done_reports, 12000)}\n\n"
        f"规则提取疑似问题：\n{json.dumps(issues, ensure_ascii=False)}\n\n"
        f"规则提取未覆盖：\n{json.dumps(uncovered, ensure_ascii=False)}\n\n"
        f"执行日志关键内容（已过滤为最终 done 与关键异常行，非原始全量日志）：\n{_trim_text(logs, 30000)}\n\n"
        f"result_json 结构化摘要（不要推断未列出的中间历史）：\n{_trim_text(result_json_summary, 30000)}"
    )
    msg = await llm.ainvoke([SystemMessage(content=system), HumanMessage(content=user)])
    data = _extract_json_object(str(getattr(msg, "content", "")))
    verdict = str(data.get("verdict") or "").strip()
    if verdict not in {"冒烟通过", "需要人工核对", "有问题"}:
        raise ValueError(f"无效 verdict: {verdict}")
    def list_field(name: str, fallback: list[str]) -> list[str]:
        value = data.get(name)
        if isinstance(value, list):
            parsed = _unique_preserve_order(str(x) for x in value)
            return parsed or fallback
        if isinstance(value, str) and value.strip():
            return _unique_preserve_order([value])
        return fallback
    ignored_scope = list_field("ignored_unrequested_scope", [])
    ignored_scope = [item for item in ignored_scope if item not in {"无", "无。", "暂无", "不适用"}]
    return {
        "verdict": verdict,
        "reason": _clean_cell(str(data.get("reason") or "")) or "LLM 基于执行日志给出结论。",
        "coverage_points": list_field("coverage_points", ["未能从 LLM 输出提取覆盖点"]),
        "problem_points": list_field("problem_points", ["未发现明确前端 bug"]),
        "manual_review_points": list_field("manual_review_points", ["无"]),
        "ignored_unrequested_scope": ignored_scope,
        "confidence": _clean_cell(str(data.get("confidence") or "")),
        "analysis_method": "llm",
    }


def build_smoke_ai_analysis_with_llm(
    *,
    report_items: list[dict[str, Any]],
    results: list[dict[str, Any]],
    issues: list[str],
    uncovered: list[str],
    custom_smoke: bool,
    requested_focus: str,
) -> dict[str, Any] | None:
    if (os.environ.get("SMOKE_AI_ANALYSIS_USE_LLM") or "true").strip().lower() in {"0", "false", "no"}:
        return None
    try:
        return asyncio.run(
            _build_smoke_ai_analysis_with_llm_async(
                report_items=report_items,
                results=results,
                issues=issues,
                uncovered=uncovered,
                custom_smoke=custom_smoke,
                requested_focus=requested_focus,
            )
        )
    except Exception as e:
        return {
            "verdict": "需要人工核对" if custom_smoke else "冒烟通过",
            "reason": f"LLM 冒烟分析失败，已退化为保守结论：{e}",
            "coverage_points": ["LLM 分析失败，请查看 HTML/日志回放。"],
            "problem_points": issues[:] if issues else ["未发现明确前端 bug"],
            "manual_review_points": uncovered[:] if uncovered else ["无"],
            "ignored_unrequested_scope": [],
            "analysis_method": "llm_error",
        }


def _health_text_indicates_ok(text: str) -> bool:
    compact = re.sub(r"\s+", "", text or "")
    ok_markers = (
        "成功访问",
        "稳定渲染",
        "渲染正常",
        "页面可达",
        "正常加载",
        "无白屏",
        "无报错",
        "无异常",
        "主体内容",
        "完整可见",
        "符合巡检标准",
    )
    bad_markers = (
        "白屏",
        "崩溃",
        "404",
        "500",
        "无法访问",
        "加载失败",
        "长期loading",
        "长时间loading",
        "报错",
        "异常",
    )
    has_ok = any(marker in compact for marker in ok_markers)
    has_bad = any(marker in compact for marker in bad_markers)
    has_negated_bad = any(marker in compact for marker in ("无白屏", "无报错", "无异常", "未发现异常", "没有异常"))
    return has_ok and (not has_bad or has_negated_bad)


def _normalize_health_report_text(text: str) -> str:
    clean = _clean_cell(text)
    if not clean:
        return ""
    if "HEALTH_REPORT" in clean:
        return _extract_health_section(clean)
    conclusion = "通过" if _health_text_indicates_ok(clean) and not _extract_smoke_issues(clean) else "部分可达"
    return "\n".join(
        [
            "HEALTH_REPORT",
            f"结论：{conclusion}",
            "入口：见执行记录中的被测入口 URL",
            "最终页面：见 quick 执行记录；Agent 已观察最终页面状态",
            f"渲染状态：{clean}",
            "异常：无" if conclusion == "通过" else f"异常：{clean}",
            "备注：由自然语言巡检完成结果自动归一化为 HEALTH_REPORT",
        ]
    )


def validate_smoke_walkthrough_attempt(
    *,
    base_validation: CompletionValidation,
    attempt_log_dir: Path,
    returncode: int,
    dry_run: bool,
) -> CompletionValidation:
    if dry_run:
        return base_validation
    try:
        from quick_completion_validate import parse_quick_attempt_outcome

        outcome = parse_quick_attempt_outcome(attempt_log_dir, returncode=returncode)
    except Exception:
        return base_validation
    done_text = _clean_cell(outcome.agent_done_text)
    if "SMOKE_REPORT" in done_text or done_text:
        issues = _extract_smoke_issues(done_text)
        uncovered = _extract_smoke_uncovered(done_text)
        missing_sections = _missing_smoke_required_sections(done_text)
        passed = returncode == 0 and bool(outcome.agent_done_success) and not missing_sections
        return CompletionValidation(
            passed=passed,
            method="smoke",
            reason=(
                "冒烟走查已完成并输出报告"
                if passed
                else (
                    f"冒烟走查报告不完整，缺少字段：{', '.join(missing_sections)}；"
                    f"疑似问题={len(issues)}，未覆盖={len(uncovered)}"
                )
            ),
            milestones_total=1,
            milestones_met=1 if done_text and not missing_sections else 0,
            missing_milestones=missing_sections or ([] if done_text else ["SMOKE_REPORT"]),
            agent_done_success=outcome.agent_done_success,
            agent_done_text=done_text,
        )
    return base_validation


def _is_smoke_early_stop_validation(validation: CompletionValidation) -> bool:
    if validation.passed or validation.method != "smoke":
        return False
    done_text = _clean_cell(validation.agent_done_text)
    if not done_text:
        return False
    fatal_markers = (
        "全站完全不可达",
        "全站不可达",
        "浏览器崩溃",
        "页面崩溃",
        "无法访问",
        "无权限",
        "登录失败",
    )
    if any(marker in done_text for marker in fatal_markers):
        return False
    return "SMOKE_REPORT" in validation.missing_milestones or bool(validation.missing_milestones)


def _smoke_early_stop_continuation_query(original_query: str, validation: CompletionValidation) -> str:
    done_text = _clean_cell(validation.agent_done_text)
    return sanitize_quick_query_for_cli(
        "\n".join(
            [
                original_query,
                "",
                "【冒烟完成度纠偏】上一轮提前结束或 SMOKE_REPORT 字段不完整，这不是可验收的任务完成。",
                f"上一轮 done 文本：{done_text[:800] or '空'}",
                "请从当前浏览器页面继续执行，不要重置任务；如果当前停在入口页或默认模块，请继续进入用户指定模块/控件并完成必要交互。",
                "页面稳定、主内容可见、等待完成、入口可达、准备进入下一步都只能作为观察结果，不能作为 done 结束理由。",
                "若局部目标仍未完成，记录为“未覆盖/需人工核对”，继续完成其它可执行检查，最后只用完整 SMOKE_REPORT 汇总。",
                "若上一轮在异步/生成中状态过早结束，请继续等待并复查，累计至少 30 秒或直到出现结果/明确错误提示；不要只按 3-5 秒等待定性。",
                "即使发现疑似问题，也应使用 done(success=true) 输出完整 SMOKE_REPORT；success=false 只用于浏览器崩溃、全站不可达等无法产出报告的情况。",
                "最终 done 文本必须以 SMOKE_REPORT 开头，并包含：结论、模块清单、模块覆盖、已检查、疑似问题、未覆盖、建议。",
            ]
        )
    )


def validate_health_check_attempt(
    *,
    base_validation: CompletionValidation,
    attempt_log_dir: Path,
    returncode: int,
    dry_run: bool,
) -> CompletionValidation:
    if dry_run:
        return base_validation
    try:
        from quick_completion_validate import parse_quick_attempt_outcome

        outcome = parse_quick_attempt_outcome(attempt_log_dir, returncode=returncode)
    except Exception:
        return base_validation
    done_text = _clean_cell(outcome.agent_done_text)
    if "HEALTH_REPORT" in done_text or done_text:
        normalized_text = _normalize_health_report_text(done_text)
        missing_sections = _missing_health_required_sections(normalized_text)
        issues = _extract_smoke_issues(normalized_text)
        passed = returncode == 0 and bool(outcome.agent_done_success) and not missing_sections and not issues
        return CompletionValidation(
            passed=passed,
            method="health",
            reason=(
                "页面巡检已完成并输出报告"
                if passed
                else f"页面巡检报告不完整，缺少字段：{', '.join(missing_sections)}；异常={len(issues)}"
            ),
            milestones_total=1,
            milestones_met=1 if normalized_text and not missing_sections else 0,
            missing_milestones=missing_sections or ([] if normalized_text else ["HEALTH_REPORT"]),
            agent_done_success=outcome.agent_done_success,
            agent_done_text=normalized_text or done_text,
        )
    return base_validation


def _html_list(items: Iterable[str]) -> str:
    clean = [html.escape(str(item)) for item in items if str(item).strip()]
    if not clean:
        clean = ["无"]
    return "<ul>" + "".join(f"<li>{item}</li>" for item in clean) + "</ul>"


def _relpath_or_abs(target: Path, start: Path) -> str:
    try:
        return os.path.relpath(target, start=start)
    except Exception:
        return str(target)


def _build_smoke_report_html(
    *,
    payload: dict[str, Any],
    title: str,
    image_base_dir: Path,
    link_base_dir: Path,
    batch_dir: Path,
) -> str:
    ai = payload.get("ai_analysis") or {}
    conclusion = html.escape(str(payload.get("conclusion") or ai.get("verdict") or "未知"))
    issue_count = int(payload.get("issue_count") or 0)
    uncovered_count = int(payload.get("uncovered_count") or 0)
    reason = html.escape(str(ai.get("reason") or ""))
    coverage_html = _html_list(ai.get("coverage_points") or [])
    problem_html = _html_list(ai.get("problem_points") or payload.get("issues") or ["未发现明确前端 bug"])
    manual_html = _html_list(ai.get("manual_review_points") or payload.get("uncovered") or ["无"])
    ignored = ai.get("ignored_unrequested_scope") or []
    ignored_html = f"<section class=\"card\"><h2>非本次重点范围</h2>{_html_list(ignored)}</section>" if ignored else ""
    status_class = "pass" if conclusion == "冒烟通过" else "bad" if conclusion == "有问题" else "warn"

    screenshot_blocks: list[str] = []
    for shot in payload.get("issue_screenshots") or []:
        rel = shot.get("relative_path")
        path = Path(shot.get("path") or "") if shot.get("path") else None
        if rel:
            src = _relpath_or_abs(batch_dir / rel, image_base_dir)
        elif path:
            src = _relpath_or_abs(path, image_base_dir)
        else:
            continue
        screenshot_blocks.append(
            "<figure>"
            f"<img src=\"{html.escape(src)}\" alt=\"截图\">"
            f"<figcaption>页面：{html.escape(str(shot.get('url') or '未知'))}<br>"
            f"说明：{html.escape(str(shot.get('reason') or '命中问题记录'))}</figcaption>"
            "</figure>"
        )
    screenshots_html = "\n".join(screenshot_blocks) if screenshot_blocks else "<p>无</p>"

    detail_rows: list[str] = []
    for item in payload.get("items") or []:
        raw_text = html.escape(str(item.get("smoke_report_text") or "未能从 quick_task_result 中提取结构化 SMOKE_REPORT，请查看原始 quick 报告。"))
        detail_rows.append(
            "<section class=\"card detail\">"
            f"<h3>{html.escape(str(item.get('case_id') or ''))} {html.escape(str(item.get('title') or ''))}</h3>"
            f"<p><b>状态：</b>{html.escape(str(item.get('status') or 'unknown'))}</p>"
            f"<p><b>result_json：</b><code>{html.escape(str(item.get('attempt_log') or '未生成'))}</code></p>"
            f"<p><b>原始 quick HTML：</b><code>{html.escape(str(item.get('html_report') or '未生成'))}</code></p>"
            f"<pre>{raw_text}</pre>"
            "</section>"
        )
    batch_md = html.escape(_relpath_or_abs(batch_dir / "smoke_walkthrough_report.md", link_base_dir))
    batch_json = html.escape(_relpath_or_abs(batch_dir / "smoke_walkthrough_report.json", link_base_dir))
    details_html = "\n".join(detail_rows)
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{html.escape(title)}</title>
  <style>
    body {{ margin: 0; font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif; color: #172033; background: #f6f8fb; }}
    main {{ max-width: 1040px; margin: 0 auto; padding: 32px 24px 48px; }}
    h1 {{ margin: 0 0 20px; font-size: 28px; }}
    h2 {{ margin-top: 28px; font-size: 18px; }}
    h3 {{ margin-top: 0; font-size: 16px; }}
    .summary {{ display: grid; grid-template-columns: repeat(4, minmax(0, 1fr)); gap: 12px; margin-bottom: 24px; }}
    .metric {{ background: #fff; border: 1px solid #e5eaf2; border-radius: 8px; padding: 14px 16px; }}
    .metric b {{ display: block; font-size: 20px; margin-top: 6px; }}
    .pass {{ color: #0f8a45; }}
    .warn {{ color: #a15c00; }}
    .bad {{ color: #b42318; }}
    .card {{ background: #fff; border: 1px solid #e5eaf2; border-radius: 8px; padding: 18px 20px; margin-top: 14px; }}
    li {{ margin: 8px 0; }}
    code {{ background: #eef2f7; padding: 2px 5px; border-radius: 4px; word-break: break-all; }}
    pre {{ white-space: pre-wrap; overflow-wrap: anywhere; background: #f3f6fa; border-radius: 6px; padding: 12px; }}
    img {{ max-width: 100%; border: 1px solid #dce3ee; border-radius: 8px; background: #fff; }}
    figcaption {{ color: #647085; margin-top: 8px; font-size: 13px; }}
    a {{ color: #1f66d1; }}
  </style>
</head>
<body>
<main>
  <h1>{html.escape(title)}</h1>
  <section class="summary">
    <div class="metric">结论<b class="{status_class}">{conclusion}</b></div>
    <div class="metric">疑似问题<b>{issue_count}</b></div>
    <div class="metric">未覆盖/风险<b>{uncovered_count}</b></div>
    <div class="metric">分析方式<b>{html.escape(str(ai.get('analysis_method') or 'unknown'))}</b></div>
  </section>
  <section class="card"><h2>AI 分析结论</h2><p>{reason}</p></section>
  <section class="card"><h2>覆盖点</h2>{coverage_html}</section>
  <section class="card"><h2>问题点</h2>{problem_html}</section>
  <section class="card"><h2>需人工核对</h2>{manual_html}</section>
  {ignored_html}
  <section class="card"><h2>截图</h2>{screenshots_html}</section>
  <section class="card"><h2>报告文件</h2><ul><li><a href="{batch_md}">Markdown 最终报告</a></li><li><a href="{batch_json}">JSON 最终报告</a></li></ul></section>
  <h2>走查明细</h2>
  {details_html}
</main>
</body>
</html>
"""


def _write_smoke_report_html_artifacts(*, payload: dict[str, Any], batch_dir: Path) -> None:
    (batch_dir / "smoke_walkthrough_report.html").write_text(
        _build_smoke_report_html(
            payload=payload,
            title="冒烟走查报告",
            image_base_dir=batch_dir,
            link_base_dir=batch_dir,
            batch_dir=batch_dir,
        ),
        encoding="utf-8",
    )
    for item in payload.get("items") or []:
        result_json = item.get("attempt_log") or ""
        if not result_json:
            continue
        attempt_dir = Path(result_json).parent
        if not attempt_dir.exists():
            continue
        (attempt_dir / "final_smoke_walkthrough_report.html").write_text(
            _build_smoke_report_html(
                payload=payload,
                title="最终冒烟走查报告",
                image_base_dir=attempt_dir,
                link_base_dir=attempt_dir,
                batch_dir=batch_dir,
            ),
            encoding="utf-8",
        )
        (attempt_dir / "final_smoke_walkthrough_report.md").write_text(
            (batch_dir / "smoke_walkthrough_report.md").read_text(encoding="utf-8"),
            encoding="utf-8",
        )
        (attempt_dir / "final_smoke_walkthrough_report.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )


def write_smoke_walkthrough_report(
    *,
    records: list[QuickRecord],
    results: list[dict[str, Any]],
    batch_dir: Path,
) -> dict[str, Any]:
    report_items: list[dict[str, Any]] = []
    all_issues: list[str] = []
    all_uncovered: list[str] = []
    for record, result in zip(records, results):
        attempts = result.get("attempts") or []
        selected_attempt = attempts[-1] if attempts else {}
        selected_text_blob = ""
        for candidate in reversed(attempts):
            validation = candidate.get("validation") or {}
            done_text = _clean_cell(validation.get("agent_done_text"))
            doc = _latest_result_doc(candidate)
            done_texts = _extract_done_texts(doc)
            text_blob = done_text or (done_texts[-1] if done_texts else "") or "\n".join(_extract_text_values(doc))
            if "SMOKE_REPORT" in text_blob:
                selected_attempt = candidate
                selected_text_blob = text_blob
                break
            if not selected_text_blob:
                selected_text_blob = text_blob
        last_attempt = selected_attempt
        text_blob = selected_text_blob
        section = _extract_smoke_section(text_blob)
        issues = _extract_smoke_issues(section or text_blob)
        uncovered = _extract_smoke_uncovered(section or text_blob)
        if result.get("status") == "failed" and not issues and not uncovered:
            issues = [f"执行未成功，需复核 quick 日志：{record.case_id}"]
        all_issues.extend(f"{record.case_id}: {issue}" for issue in issues)
        all_uncovered.extend(f"{record.case_id}: {item}" for item in uncovered)
        report_items.append(
            {
                "case_id": record.case_id,
                "title": record.title,
                "status": result.get("status"),
                "issues": issues,
                "uncovered": uncovered,
                "smoke_report_text": section,
                "attempt_log": last_attempt.get("result_json") or "",
                "html_report": last_attempt.get("report_html") or "",
                "original_case_path": result.get("original_case_path") or "",
            }
        )

    custom_smoke = any(bool((record.raw or {}).get("custom_smoke")) for record in records)
    requested_focus = "；".join(
        str((record.raw or {}).get("smoke_focus") or "").strip()
        for record in records
        if str((record.raw or {}).get("smoke_focus") or "").strip()
    )
    issue_screenshots = _save_issue_screenshots(results, batch_dir, has_final_issues=bool(all_issues))

    ai_analysis = build_smoke_ai_analysis(
        report_items=report_items,
        results=results,
        issues=all_issues,
        uncovered=all_uncovered,
        custom_smoke=custom_smoke,
        issue_screenshots=issue_screenshots,
        requested_focus=requested_focus,
    )
    conclusion = ai_analysis["verdict"]
    display_issues = [
        str(item)
        for item in (ai_analysis.get("problem_points") or [])
        if str(item).strip() and "未发现明确前端 bug" not in str(item)
    ]
    display_uncovered = [
        str(item)
        for item in (ai_analysis.get("manual_review_points") or [])
        if str(item).strip() and str(item).strip() not in {"无", "无。", "暂无", "不适用"}
    ]
    if not display_issues:
        issue_screenshots = []
        ai_analysis["issue_screenshots"] = []

    md_lines = [
        "# 冒烟走查报告",
        "",
        f"- 结论：{conclusion}",
        f"- 记录数：{len(report_items)}",
        f"- 疑似问题数：{len(display_issues)}",
        f"- 未覆盖/风险说明数：{len(display_uncovered)}",
        "",
        "## 疑似问题",
    ]
    if display_issues:
        md_lines.extend(f"- {issue}" for issue in display_issues)
    else:
        md_lines.append("- 无")
    md_lines.extend(["", "## 未覆盖 / 风险说明"])
    if display_uncovered:
        md_lines.extend(f"- {item}" for item in display_uncovered)
    else:
        md_lines.append("- 无")
    md_lines.extend(
        [
            "",
            "## AI 分析结论",
            f"- 明确结论：{ai_analysis['verdict']}",
            f"- 判断依据：{ai_analysis['reason']}",
            "- 覆盖点：",
        ]
    )
    md_lines.extend(f"  - {item}" for item in ai_analysis["coverage_points"])
    md_lines.append("- 问题点：")
    md_lines.extend(f"  - {item}" for item in ai_analysis["problem_points"])
    md_lines.append("- 需人工核对：")
    md_lines.extend(f"  - {item}" for item in ai_analysis["manual_review_points"])
    ignored_scope = ai_analysis.get("ignored_unrequested_scope") or []
    if ignored_scope:
        md_lines.append("- 非本次重点范围：")
        md_lines.extend(f"  - {item}" for item in ignored_scope)
    if issue_screenshots:
        md_lines.extend(["", "## 问题截图"])
        for shot in issue_screenshots:
            md_lines.extend(
                [
                    f"![问题截图 {shot.get('history_index', '')}]({shot['relative_path']})",
                    f"- 页面：{shot.get('url') or '未知'}",
                    f"- 说明：{shot.get('reason') or '命中问题记录'}",
                    "",
                ]
            )
    md_lines.extend(["", "## 走查明细"])
    for item in report_items:
        md_lines.extend(
            [
                f"### {item['case_id']} {item['title']}",
                "",
                f"- 状态：{item['status']}",
                f"- result_json：`{item['attempt_log'] or '未生成'}`",
                f"- html_report：`{item['html_report'] or '未生成'}`",
                "",
                item["smoke_report_text"] or "未能从 quick_task_result 中提取结构化 SMOKE_REPORT，请查看 HTML 报告。",
                "",
            ]
        )
    payload = {
        "conclusion": conclusion,
        "issue_count": len(display_issues),
        "uncovered_count": len(display_uncovered),
        "issues": display_issues,
        "uncovered": display_uncovered,
        "raw_issue_count": len(all_issues),
        "raw_uncovered_count": len(all_uncovered),
        "raw_issues": all_issues,
        "raw_uncovered": all_uncovered,
        "ai_analysis": ai_analysis,
        "issue_screenshots": issue_screenshots,
        "items": report_items,
    }
    (batch_dir / "smoke_walkthrough_report.md").write_text("\n".join(md_lines).strip() + "\n", encoding="utf-8")
    (batch_dir / "smoke_walkthrough_report.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    _write_smoke_report_html_artifacts(payload=payload, batch_dir=batch_dir)
    ai_md_lines = [
        "# 冒烟 AI 分析结论",
        "",
        f"- 明确结论：{ai_analysis['verdict']}",
        f"- 判断依据：{ai_analysis['reason']}",
        "",
        "## 覆盖点",
        *(f"- {item}" for item in ai_analysis["coverage_points"]),
        "",
        "## 问题点",
        *(f"- {item}" for item in ai_analysis["problem_points"]),
        "",
        "## 需人工核对",
        *(f"- {item}" for item in ai_analysis["manual_review_points"]),
    ]
    ignored_scope = ai_analysis.get("ignored_unrequested_scope") or []
    if ignored_scope:
        ai_md_lines.extend(["", "## 非本次重点范围", *(f"- {item}" for item in ignored_scope)])
    if issue_screenshots:
        ai_md_lines.extend(["", "## 问题截图"])
        for shot in issue_screenshots:
            ai_md_lines.extend(
                [
                    f"![问题截图 {shot.get('history_index', '')}]({shot['relative_path']})",
                    f"- 页面：{shot.get('url') or '未知'}",
                    f"- 说明：{shot.get('reason') or '命中问题记录'}",
                    "",
                ]
            )
    (batch_dir / "smoke_ai_analysis.md").write_text("\n".join(ai_md_lines).strip() + "\n", encoding="utf-8")
    (batch_dir / "smoke_ai_analysis.json").write_text(
        json.dumps(ai_analysis, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return payload


def write_health_check_report(
    *,
    records: list[QuickRecord],
    results: list[dict[str, Any]],
    batch_dir: Path,
) -> dict[str, Any]:
    report_items: list[dict[str, Any]] = []
    all_issues: list[str] = []
    for record, result in zip(records, results):
        attempts = result.get("attempts") or []
        selected_attempt = attempts[-1] if attempts else {}
        selected_text_blob = ""
        for candidate in reversed(attempts):
            validation = candidate.get("validation") or {}
            done_text = _clean_cell(validation.get("agent_done_text"))
            doc = _latest_result_doc(candidate)
            done_texts = _extract_done_texts(doc)
            text_blob = done_text or (done_texts[-1] if done_texts else "") or "\n".join(_extract_text_values(doc))
            if "HEALTH_REPORT" in text_blob:
                selected_attempt = candidate
                selected_text_blob = text_blob
                break
            if not selected_text_blob:
                selected_text_blob = text_blob
        section = _normalize_health_report_text(selected_text_blob) or _extract_health_section(selected_text_blob)
        issues = _extract_smoke_issues(section or selected_text_blob)
        if result.get("status") == "failed" and _health_text_indicates_ok(section or selected_text_blob) and not issues:
            result["status"] = "success"
        if result.get("status") == "failed" and not issues:
            issues = [f"页面巡检执行未成功，需复核 quick 日志：{record.case_id}"]
        all_issues.extend(f"{record.case_id}: {issue}" for issue in issues)
        report_items.append(
            {
                "case_id": record.case_id,
                "title": record.title,
                "status": result.get("status"),
                "issues": issues,
                "health_report_text": section,
                "attempt_log": selected_attempt.get("result_json") or "",
                "html_report": selected_attempt.get("report_html") or "",
                "original_case_path": result.get("original_case_path") or "",
            }
        )

    conclusion = "通过" if not all_issues else "异常"
    md_lines = [
        "# 页面巡检报告",
        "",
        f"- 结论：{conclusion}",
        f"- 记录数：{len(report_items)}",
        f"- 异常数：{len(all_issues)}",
        "",
        "## 异常",
    ]
    if all_issues:
        md_lines.extend(f"- {issue}" for issue in all_issues)
    else:
        md_lines.append("- 无")
    md_lines.extend(["", "## 巡检明细"])
    for item in report_items:
        md_lines.extend(
            [
                f"### {item['case_id']} {item['title']}",
                "",
                f"- 状态：{item['status']}",
                f"- result_json：`{item['attempt_log'] or '未生成'}`",
                f"- html_report：`{item['html_report'] or '未生成'}`",
                "",
                item["health_report_text"] or "未能从 quick_task_result 中提取结构化 HEALTH_REPORT，请查看 HTML 报告。",
                "",
            ]
        )
    payload = {
        "conclusion": conclusion,
        "issue_count": len(all_issues),
        "issues": all_issues,
        "items": report_items,
    }
    (batch_dir / "health_check_report.md").write_text("\n".join(md_lines).strip() + "\n", encoding="utf-8")
    (batch_dir / "health_check_report.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return payload


def original_case_markdown(record: QuickRecord) -> str:
    """Render the source/generated test case before it is rewritten into quick query."""
    raw = record.raw or {}
    lines = [
        f"# {record.case_id} {record.title}",
        "",
        f"- 来源：{record.source}",
        f"- 是否执行：{'是' if record.should_run else '否'}",
    ]
    if record.skip_reason:
        lines.append(f"- 跳过原因：{record.skip_reason}")
    lines.append("")
    lines.append("## 原始用例")
    if raw:
        preferred = [
            ("title", "用例标题"),
            ("module", "模块"),
            ("module1", "一级模块"),
            ("module2", "二级模块"),
            ("priority", "优先级"),
            ("smoke", "是否冒烟"),
            ("case_type", "类型"),
            ("attribute", "属性"),
            ("test_point", "测试点"),
            ("tp", "测试点"),
            ("precondition", "前置条件"),
            ("steps", "测试步骤"),
            ("expected", "预期结果"),
            ("ui_auto", "是否 UI 自动化"),
            ("interface_auto", "是否接口自动化"),
            ("fixture_required", "是否需要前置数据"),
            ("fixture", "前置数据说明"),
            ("entry_hint", "建议入口"),
            ("navigation", "导航路径"),
            ("automation_note", "自动化执行备注"),
            ("description", "描述"),
            ("source_chunk", "来源片段"),
        ]
        emitted: set[str] = set()
        for key, label in preferred:
            if key not in raw:
                continue
            value = _normalize_case_value(raw.get(key))
            if value:
                lines.extend([f"### {label}", value, ""])
                emitted.add(key)
        for key in sorted(k for k in raw if k not in emitted):
            value = _normalize_case_value(raw.get(key))
            if value:
                lines.extend([f"### {key}", value, ""])
    else:
        lines.append("该记录来自单条 query，未经过用例集合或 PRD/测分文档拆解。")
        lines.append("")
    lines.extend(["## 改写后的 quick query", record.query, ""])
    return "\n".join(lines).strip() + "\n"


def write_original_case_artifacts(records: list[QuickRecord], *, batch_dir: Path, out_dir: Path) -> None:
    batch_dir.mkdir(parents=True, exist_ok=True)
    cases_dir = batch_dir / "original_cases"
    cases_dir.mkdir(parents=True, exist_ok=True)
    index_lines = ["# 原始用例总览", ""]
    json_items: list[dict[str, Any]] = []
    for record in records:
        md = original_case_markdown(record)
        case_path = cases_dir / f"{record.case_id}.md"
        case_path.write_text(md, encoding="utf-8")
        log_case_dir = out_dir / "logs" / record.case_id
        log_case_dir.mkdir(parents=True, exist_ok=True)
        (log_case_dir / "original_case.md").write_text(md, encoding="utf-8")
        index_lines.extend(
            [
                f"## {record.case_id} {record.title}",
                "",
                f"- 执行日志内原始用例：`logs/{record.case_id}/original_case.md`",
                f"- 批量原始用例文件：`batch/original_cases/{record.case_id}.md`",
                "",
                md,
            ]
        )
        json_items.append(
            {
                "case_id": record.case_id,
                "title": record.title,
                "source": record.source,
                "should_run": record.should_run,
                "skip_reason": record.skip_reason,
                "raw": record.raw,
                "query": record.query,
                "original_case_md": str(case_path),
                "log_original_case_md": str(log_case_dir / "original_case.md"),
            }
        )
    (batch_dir / "original_cases.md").write_text("\n".join(index_lines).strip() + "\n", encoding="utf-8")
    (batch_dir / "original_cases.json").write_text(json.dumps(json_items, ensure_ascii=False, indent=2), encoding="utf-8")


def _slug_case_id(idx: int, title: str) -> str:
    text = re.sub(r"[^\w\u4e00-\u9fff.-]+", "_", title.strip(), flags=re.UNICODE).strip("_")
    if not text:
        text = "case"
    return f"{idx:03d}_{text[:48]}"


def _column_map(headers: list[str]) -> dict[str, int]:
    norm = {h.strip().lower(): i for i, h in enumerate(headers)}
    out: dict[str, int] = {}
    for key, aliases in CASE_COLUMNS.items():
        for alias in aliases:
            idx = norm.get(alias.strip().lower())
            if idx is not None:
                out[key] = idx
                break
    return out


def _row_to_record(idx: int, row: dict[str, str], source: str) -> QuickRecord:
    title = _clean_cell(row.get("title") or row.get("description") or f"测试记录 {idx}")
    ui_auto = _clean_cell(row.get("ui_auto"))
    should_run = "不可" not in ui_auto if ui_auto else True
    skip_reason = "marked_non_ui_automation" if not should_run else ""
    parts = [
        f"用例标题：{title}",
        f"一级模块：{_clean_cell(row.get('module1'))}" if row.get("module1") else "",
        f"二级模块：{_clean_cell(row.get('module2'))}" if row.get("module2") else "",
        f"等级：{_clean_cell(row.get('priority'))}" if row.get("priority") else "",
        f"类型：{_clean_cell(row.get('case_type'))}" if row.get("case_type") else "",
        f"前置条件：{_clean_cell(row.get('precondition'))}" if row.get("precondition") else "",
        f"测试步骤：{_clean_cell(row.get('steps'))}" if row.get("steps") else "",
        f"预期结果：{_clean_cell(row.get('expected'))}" if row.get("expected") else "",
        f"补充描述：{_clean_cell(row.get('description'))}" if row.get("description") else "",
    ]
    query = "。".join(p for p in parts if p)
    query += "。请按该测试记录执行 UI 快速验证，并在完成后判断预期结果是否满足。"
    return QuickRecord(
        case_id=_slug_case_id(idx, title),
        title=title,
        query=query,
        source=source,
        should_run=should_run,
        skip_reason=skip_reason,
        raw={k: v for k, v in row.items() if v},
    )


def _parse_markdown_table(path: Path) -> list[QuickRecord]:
    text = path.read_text(encoding="utf-8")
    marker = "## 测试用例总表"
    if marker in text:
        text = text.split(marker, 1)[1]
    lines = [ln.rstrip() for ln in text.splitlines()]
    table_lines: list[str] = []
    in_table = False
    for ln in lines:
        if ln.strip().startswith("|"):
            table_lines.append(ln)
            in_table = True
        elif in_table:
            break
    if len(table_lines) < 3:
        return []

    def split_row(ln: str) -> list[str]:
        body = ln.strip().strip("|")
        return [_clean_cell(c) for c in body.split("|")]

    headers = split_row(table_lines[0])
    cmap = _column_map(headers)
    if len(cmap) < 2 or not ({"title", "steps", "expected", "description"} & set(cmap)):
        return []
    records: list[QuickRecord] = []
    for raw_idx, ln in enumerate(table_lines[2:], start=1):
        cells = split_row(ln)
        if not any(cells):
            continue
        row = {key: cells[pos] if pos < len(cells) else "" for key, pos in cmap.items()}
        records.append(_row_to_record(raw_idx, row, str(path)))
    return records


def _parse_spreadsheet(path: Path) -> list[QuickRecord]:
    if path.suffix.lower() in (".csv", ".tsv"):
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f, delimiter=delimiter))
        return _records_from_rows(rows, str(path))

    try:
        from openpyxl import load_workbook
    except Exception as e:  # pragma: no cover - depends on optional local package
        raise RuntimeError("解析 xlsx 需要 openpyxl，请在同一 Python 环境安装 openpyxl") from e

    wb = load_workbook(path, read_only=True, data_only=True)
    best_rows: list[list[Any]] = []
    best_score = -1
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        headers = [_clean_cell(c) for c in rows[0]]
        score = len(_column_map(headers))
        if score > best_score:
            best_score = score
            best_rows = rows
    return _records_from_rows(best_rows, str(path))


def _records_from_rows(rows: list[list[Any]], source: str) -> list[QuickRecord]:
    if len(rows) < 2:
        return []
    headers = [_clean_cell(c) for c in rows[0]]
    cmap = _column_map(headers)
    if len(cmap) < 2 or not ({"title", "steps", "expected", "description"} & set(cmap)):
        return []
    records: list[QuickRecord] = []
    for idx, cells in enumerate(rows[1:], start=1):
        if not any(_clean_cell(c) for c in cells):
            continue
        row = {key: _clean_cell(cells[pos]) if pos < len(cells) else "" for key, pos in cmap.items()}
        records.append(_row_to_record(idx, row, source))
    return records


def load_case_records(input_path: Path) -> list[QuickRecord]:
    suffix = input_path.suffix.lower()
    if suffix == ".md":
        return _parse_markdown_table(input_path)
    if suffix in (".xlsx", ".csv", ".tsv"):
        return _parse_spreadsheet(input_path)
    return []


def _read_xlsx_as_text(path: Path, *, max_rows: int = 240, max_cols: int = 16) -> str:
    try:
        from openpyxl import load_workbook
    except Exception as e:  # pragma: no cover - depends on optional local package
        raise RuntimeError("解析 xlsx 文档需要 openpyxl，请在同一 Python 环境安装 openpyxl") from e
    wb = load_workbook(path, read_only=True, data_only=True)
    chunks: list[str] = []
    for ws in wb.worksheets[:6]:
        chunks.append(f"# Sheet: {ws.title}")
        count = 0
        for row in ws.iter_rows(values_only=True):
            cells = [_clean_cell(c) for c in list(row)[:max_cols]]
            if any(cells):
                chunks.append(" | ".join(cells))
                count += 1
            if count >= max_rows:
                chunks.append("...（已截断）")
                break
    return "\n".join(chunks).strip()


def read_source_document(path: Path, *, limit: int = 80000) -> str:
    suffix = path.suffix.lower()
    if path.is_dir():
        summary = build_context_summary([path], max_files=32, max_chars_per_file=2400)
        return summary[:limit]
    if suffix in DOC_TEXT_SUFFIXES:
        return _read_small_text(path, limit)
    if suffix == ".xlsx":
        return _read_xlsx_as_text(path)[:limit]
    if suffix == ".docx":
        try:
            from docx import Document
        except Exception as e:  # pragma: no cover - optional dependency
            raise RuntimeError("解析 docx 文档需要 python-docx，请在同一 Python 环境安装 python-docx") from e
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs if p.text).strip()[:limit]
    raise ValueError(f"不支持的文档输入类型: {path.suffix}")


def detect_input_kind(input_path: Path, source_text: str, parsed_cases: list[QuickRecord]) -> str:
    if parsed_cases:
        return "cases"
    blob = f"{input_path.name}\n{source_text[:6000]}".lower()
    if any(k.lower() in blob for k in DOC_KEYWORDS_ANALYSIS):
        return "analysis"
    if any(k.lower() in blob for k in DOC_KEYWORDS_REQUIREMENT):
        return "requirement"
    return "unknown"


def _get_doc_llm() -> Any:
    from browser_use.agent.chat_models import build_models

    default_model = "qwen3.7-plus"

    def env_model(key: str, fallback: str) -> str:
        value = (os.environ.get(key) or "").strip()
        return value if value else fallback

    dm = env_model("BU_PAGE_EXTRACTION_LLM", default_model)
    actor = env_model("BU_ACTOR_MODEL", dm)
    planner = env_model("BU_PLANNER_MODEL", dm)
    asserter = env_model("BU_ASSERTER_MODEL", dm)
    page = env_model("BU_PAGE_EXTRACTION_LLM", dm)
    locating = env_model("BU_LOCATING_LLM", page)
    coding = env_model("BU_CODING_LLM", page)
    _, _, _, aux = build_models(
        actor=actor,
        planner=planner,
        asserter=asserter,
        page_extraction_llm=page,
        locating_llm=locating,
        coding_llm=coding,
    )
    llm = aux.get("page_extraction_llm")
    if llm is None:
        raise RuntimeError("page_extraction_llm 不可用，请检查 BU_OPENAI_* 或兼容的模型配置")
    return llm


def _extract_json_array(text: str) -> list[dict[str, Any]]:
    raw = (text or "").strip()
    raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.I)
    raw = re.sub(r"\s*```$", "", raw)
    start = raw.find("[")
    end = raw.rfind("]")
    if start < 0 or end < start:
        raise ValueError("模型未返回 JSON 数组")
    data = json.loads(raw[start : end + 1])
    if not isinstance(data, list):
        raise ValueError("模型返回内容不是 JSON 数组")
    return [x for x in data if isinstance(x, dict)]


def _split_doc_for_generation(source_text: str, *, max_chars: int) -> list[tuple[str, str]]:
    """Split large requirement docs by markdown headings before asking the LLM."""
    text = (source_text or "").strip()
    if not text:
        return []
    max_chars = max(4000, max_chars)
    sections: list[tuple[str, str]] = []
    current_title = "文档开头"
    current_lines: list[str] = []
    heading_re = re.compile(r"^(#{1,4})\s+(.+?)\s*$")
    for line in text.splitlines():
        m = heading_re.match(line.strip())
        if m and current_lines:
            sections.append((current_title, "\n".join(current_lines).strip()))
            current_title = m.group(2).strip()
            current_lines = [line]
        else:
            if m:
                current_title = m.group(2).strip()
            current_lines.append(line)
    if current_lines:
        sections.append((current_title, "\n".join(current_lines).strip()))

    chunks: list[tuple[str, str]] = []
    buf_title = ""
    buf: list[str] = []
    for title, body in sections:
        if not body:
            continue
        if len(body) > max_chars:
            if buf:
                chunks.append((buf_title or "文档片段", "\n\n".join(buf).strip()))
                buf_title = ""
                buf = []
            for idx in range(0, len(body), max_chars):
                chunks.append((f"{title} part {idx // max_chars + 1}", body[idx : idx + max_chars]))
            continue
        candidate = "\n\n".join(buf + [body])
        if buf and len(candidate) > max_chars:
            chunks.append((buf_title or "文档片段", "\n\n".join(buf).strip()))
            buf = [body]
            buf_title = title
        else:
            buf.append(body)
            buf_title = f"{buf_title} / {title}" if buf_title else title
    if buf:
        chunks.append((buf_title or "文档片段", "\n\n".join(buf).strip()))
    return chunks or [("文档全文", text[:max_chars])]


def _dedupe_doc_items(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    seen: set[str] = set()
    out: list[dict[str, Any]] = []
    for item in items:
        title = _clean_cell(item.get("title") or item.get("用例标题"))
        module = _clean_cell(item.get("module") or item.get("模块"))
        steps = item.get("steps") or item.get("测试步骤") or ""
        if isinstance(steps, list):
            steps_text = " ".join(_clean_cell(s) for s in steps[:3])
        else:
            steps_text = _clean_cell(steps)
        key = re.sub(r"\s+", "", f"{module}|{title}|{steps_text[:80]}").lower()
        if not key or key in seen:
            continue
        seen.add(key)
        out.append(item)
    return out


DOC_META_PATTERNS = (
    "变更记录",
    "修订记录",
    "需求背景",
    "产品目标",
    "长期目标",
    "常态化运营",
    "需求列表",
    "文档说明",
    "版本记录",
)
UI_ACTION_PATTERNS = (
    "点击",
    "输入",
    "选择",
    "勾选",
    "上传",
    "打开",
    "进入",
    "切换",
    "搜索",
    "筛选",
    "保存",
    "提交",
    "删除",
    "编辑",
    "新增",
    "创建",
    "查看",
    "展开",
    "关闭",
    "复制",
    "点赞",
    "点踩",
    "悬浮",
    "滚动",
)
FIXTURE_PATTERNS = (
    "已存在",
    "已有",
    "已创建",
    "已发布",
    "预置",
    "准备",
    "至少",
    "同名",
    "原有",
    "绑定",
    "关联",
    "已上传",
    "处理成功",
)
UPLOAD_PATTERNS = ("上传", "选择文件", "本地文件", "PDF", "JPG", "PNG", "HTML", "表格文件", "含图文档", "大文件", "损坏文件")
ASYNC_PATTERNS = ("等待解析", "解析完成", "处理完成", "处理中", "任务状态", "重试", "异步", "导入成功")


def _bool_like(value: Any, default: bool = True) -> bool:
    if isinstance(value, bool):
        return value
    text = _clean_cell(value).lower()
    if not text:
        return default
    if text in {"false", "no", "0", "否", "不可", "不可ui自动化", "不可 ui 自动化"}:
        return False
    if "不可" in text:
        return False
    return True


def _has_any(text: str, patterns: Iterable[str]) -> bool:
    return any(p.lower() in text.lower() for p in patterns)


def _normalize_doc_item(item: dict[str, Any], *, input_kind: str, chunk_title: str) -> dict[str, Any]:
    """Harden LLM-generated PRD/test-analysis cases before turning them into quick records."""
    normalized = dict(item)
    title = _clean_cell(normalized.get("title") or normalized.get("用例标题"))
    module = _clean_cell(normalized.get("module") or normalized.get("模块"))
    module1 = _clean_cell(normalized.get("module1") or normalized.get("一级模块"))
    module2 = _clean_cell(normalized.get("module2") or normalized.get("二级模块"))
    test_point = _clean_cell(normalized.get("test_point") or normalized.get("测试点") or normalized.get("tp"))
    precondition = _clean_cell(normalized.get("precondition") or normalized.get("前置条件"))
    expected = _clean_cell(normalized.get("expected") or normalized.get("预期结果"))
    steps = normalized.get("steps") or normalized.get("测试步骤") or []
    if isinstance(steps, list):
        steps_text = " ".join(_clean_cell(s) for s in steps if _clean_cell(s))
    else:
        steps_text = _clean_cell(steps)
    case_blob = "\n".join([title, module, module1, module2, test_point, precondition, steps_text, expected])
    blob = "\n".join([case_blob, chunk_title])

    ui_auto = _bool_like(normalized.get("ui_auto", True), default=True)
    reasons: list[str] = []
    if input_kind == "requirement" and _has_any(case_blob, DOC_META_PATTERNS):
        # Requirement metadata/background may inform coverage, but it is not a product UI target.
        ui_auto = False
        reasons.append("需求文档元信息/背景目标，不作为产品 UI quick 用例执行")
    if not _has_any(steps_text, UI_ACTION_PATTERNS):
        ui_auto = False
        reasons.append("测试步骤缺少明确 UI 原子操作")
    if re.search(r"(变更记录|需求背景|产品目标).{0,30}(表格|章节|标题|文案|展示)", case_blob):
        ui_auto = False
        reasons.append("把 PRD 文档章节当作页面内容验证")

    fixture_required = _bool_like(normalized.get("fixture_required"), default=False)
    fixture = _clean_cell(normalized.get("fixture") or normalized.get("前置数据说明"))
    if _has_any(blob, FIXTURE_PATTERNS) or _has_any(blob, UPLOAD_PATTERNS) or _has_any(blob, ASYNC_PATTERNS):
        fixture_required = True
        if not fixture:
            fixture = "需要测试数据/文件/业务状态满足前置条件；若未配置 fixture，应先造数或跳过该用例"

    if reasons:
        note = _clean_cell(normalized.get("automation_note"))
        normalized["automation_note"] = "；".join([p for p in (note, "；".join(reasons)) if p])
    normalized["ui_auto"] = ui_auto
    normalized["fixture_required"] = fixture_required
    if fixture:
        normalized["fixture"] = fixture
    normalized.setdefault("attribute", "功能")
    normalized.setdefault("source", "需求-功能" if input_kind == "requirement" else "测分-功能")
    if not normalized.get("module1") and module:
        if "-" in module:
            first, second = module.split("-", 1)
            normalized["module1"] = first.strip()
            normalized["module2"] = normalized.get("module2") or second.strip()
        else:
            normalized["module1"] = module
    return normalized


def _doc_item_to_record(idx: int, item: dict[str, Any], *, source: str, task_prefix: str) -> QuickRecord:
    title = _clean_cell(item.get("title") or item.get("用例标题") or f"文档拆解用例 {idx}")
    steps = item.get("steps") or item.get("测试步骤") or []
    if isinstance(steps, list):
        steps_text = "；".join(_clean_cell(s) for s in steps if _clean_cell(s))
    else:
        steps_text = _clean_cell(steps)
    expected = _clean_cell(item.get("expected") or item.get("预期结果"))
    precondition = _clean_cell(item.get("precondition") or item.get("前置条件"))
    priority = _clean_cell(item.get("priority") or item.get("优先级"))
    module = _clean_cell(item.get("module") or item.get("模块"))
    module1 = _clean_cell(item.get("module1") or item.get("一级模块"))
    module2 = _clean_cell(item.get("module2") or item.get("二级模块"))
    test_point = _clean_cell(item.get("test_point") or item.get("测试点") or item.get("tp"))
    ui_auto = _bool_like(item.get("ui_auto", item.get("是否UI自动化", True)), default=True)
    skip_reason = "" if ui_auto else "generated_non_ui_automation"
    fixture_required = _bool_like(item.get("fixture_required"), default=False)
    fixture = _clean_cell(item.get("fixture") or item.get("前置数据说明"))
    entry_hint = _clean_cell(item.get("entry_hint") or item.get("建议入口"))
    navigation = _clean_cell(item.get("navigation") or item.get("导航路径"))
    automation_note = _clean_cell(item.get("automation_note") or item.get("自动化执行备注"))
    smoke = _clean_cell(item.get("smoke") or item.get("是否冒烟"))
    case_type = _clean_cell(item.get("case_type") or item.get("类型"))
    record_prefix = (task_prefix or "").strip()
    parts = [
        record_prefix,
        f"用例标题：{title}",
        f"一级模块：{module1}" if module1 else "",
        f"二级模块：{module2}" if module2 else "",
        f"模块：{module}" if module and not (module1 or module2) else "",
        f"优先级：{priority}" if priority else "",
        f"是否冒烟：{smoke}" if smoke else "",
        f"类型：{case_type}" if case_type else "",
        f"覆盖测试点：{test_point}" if test_point else "",
        f"前置条件：{precondition}" if precondition else "",
        f"建议入口：{entry_hint}" if entry_hint else "",
        f"导航路径：{navigation}" if navigation else "",
        f"前置数据：{fixture}" if fixture_required and fixture else "",
        f"测试步骤：{steps_text}" if steps_text else "",
        f"预期结果：{expected}" if expected else "",
        f"自动化执行备注：{automation_note}" if automation_note and not ui_auto else "",
    ]
    query = "。".join(p for p in parts if p)
    query += "。请按该测试用例执行 UI 快速验证，并在完成后判断预期结果是否满足。"
    return QuickRecord(
        case_id=_slug_case_id(idx, title),
        title=title,
        query=query,
        source=source,
        should_run=ui_auto,
        skip_reason=skip_reason,
        raw=item,
    )


async def _generate_records_from_doc_async(
    *,
    input_path: Path,
    input_kind: str,
    source_text: str,
    context_summary: str,
    task_prefix: str,
    max_generated_cases: int,
    doc_chunk_chars: int,
    cases_per_doc_chunk: int,
) -> list[QuickRecord]:
    try:
        from langchain_core.messages import HumanMessage, SystemMessage
    except ImportError as e:
        raise RuntimeError("文档拆解需要 langchain_core，请使用安装了 browser-use 3.1.6 的 Python 环境") from e

    llm = _get_doc_llm()
    kind_label = "需求文档" if input_kind == "requirement" else "测分/测试分析文档"
    global_limit_text = "不限制总条数" if max_generated_cases <= 0 else f"总数最多 {max_generated_cases} 条"
    per_chunk_limit_text = "不限制本片段条数" if cases_per_doc_chunk <= 0 else f"本片段最多 {cases_per_doc_chunk} 条"
    system = (
        "你是资深 UI 自动化测试设计专家。请把输入文档拆解为可逐条执行的 UI quick 测试用例。\n"
        "先在内部完成“可测点识别→非测试点过滤→用例编写→执行友好性自检”，最终只输出 JSON 数组，不要 markdown，不要解释。\n"
        "数组元素字段固定为：title, module1, module2, priority, smoke, case_type, attribute, test_point, "
        "precondition, entry_hint, navigation, steps, expected, ui_auto, interface_auto, fixture_required, fixture, automation_note, source。\n"
        "steps 必须是字符串数组；ui_auto/interface_auto/fixture_required 只能是 true 或 false。\n"
        "字段取值：priority=P0/P1/P2；smoke=是/否；case_type=核心/常规/异常；attribute=功能/非功能；source=需求-功能/需求-非功能/测分-功能/测分-非功能。\n"
        "过滤硬规则：需求背景、产品目标、长期目标、常态化运营、变更记录、修订记录、需求列表链接、图示占位、过程管理、无可观察系统对象的段落，不生成可执行 UI 用例；"
        "若必须留痕则输出 ui_auto=false，并在 automation_note 说明“非产品 UI 验证点”。禁止把 PRD 文档章节当成被测页面章节验证。\n"
        "覆盖规则：默认全量覆盖输入片段中的有效功能点、表格行、并列子点、枚举值、配置规则、边界校验、列表筛选、状态流转和关键交互；不要只挑 3-5 条代表性场景。\n"
        "用例粒度：正向成功、异常/边界、不同枚举/文件类型/状态流转原则上拆成独立用例；单条用例必须单目标、线性、可由一句 quick query 执行。\n"
        "UI 执行友好规则：precondition 写清登录角色、入口、目标页、必要数据；entry_hint 写建议路由或入口模块；navigation 写从入口到目标页的可见菜单/Tab 路径；"
        "steps 必须是原子动作数组，一步只点击/输入/选择/等待/勾选/上传/查看一种动作，使用页面可见文案、输入框标签、Tab 名称；禁止“按需求验证”“功能正常”“若需要则”等模糊或分支表述。\n"
        "预期结果必须与步骤可对应，使用可观察断言：toast 文案、列表列值、详情字段、Tab 高亮、弹窗内容、URL/页面标题、无错误页；最后一条必须能判定成功/失败。\n"
        "fixture 规则：凡需要已存在标签/文件/知识库/已发布知识/处理成功任务/上传样例文件/特定权限或状态，fixture_required=true，并在 fixture 写清缺少哪些测试数据；"
        "没有 fixture 时不要把不可满足的前置伪装成已满足。纯接口、后端、性能、安全越权类场景 ui_auto=false。"
    )
    chunks = _split_doc_for_generation(source_text, max_chars=doc_chunk_chars)
    all_items: list[dict[str, Any]] = []
    for chunk_idx, (chunk_title, chunk_text) in enumerate(chunks, start=1):
        user = (
            f"输入类型：{kind_label}\n"
            f"来源文件：{input_path}\n"
            f"生成数量要求：{global_limit_text}；{per_chunk_limit_text}\n"
            f"当前片段：{chunk_idx}/{len(chunks)} - {chunk_title}\n"
            f"任务公共前缀（入口 URL、登录态、租户等，若为空可忽略）：\n{task_prefix or '(无)'}\n\n"
            f"可选辅助背景（产品白皮书、代码库、图谱目录等，若为空可忽略）：\n{context_summary[:12000] or '(无)'}\n\n"
            "拆解要求：\n"
            "- 本片段每个明确功能模块、需求明细、表格行或并列子点至少生成 1 条用例记录；只有产品 UI 可操作且可观察的记录 ui_auto=true。\n"
            "- 有新增/编辑/删除/筛选/批量/状态/失败重试/权限/必填/超长/非法值/空列表时，应分别成用例。\n"
            "- 正向成功与异常/边界场景分开成用例，不要合并成一条带条件分支的用例。\n"
            "- 如果某项不是 UI 可验证，也输出 ui_auto=false 记录，便于报告说明未执行原因；但不要输出 PRD 元信息型伪用例。\n"
            "- 需求文档场景下，重点从“需求说明/需求详情/功能模块/验收标准”提取，不要从“变更记录/需求背景/产品目标”生成 quick 执行用例。\n\n"
            f"待拆解文档片段：\n{chunk_text}"
        )
        msg = await llm.ainvoke([SystemMessage(content=system), HumanMessage(content=user)])
        content = getattr(msg, "content", "")
        chunk_items = _extract_json_array(str(content))
        if cases_per_doc_chunk > 0:
            chunk_items = chunk_items[:cases_per_doc_chunk]
        for item in chunk_items:
            item.setdefault("source_chunk", chunk_title)
            item = _normalize_doc_item(item, input_kind=input_kind, chunk_title=chunk_title)
            all_items.append(item)

    items = _dedupe_doc_items(all_items)
    if max_generated_cases > 0:
        items = items[:max_generated_cases]
    return [
        _doc_item_to_record(i, item, source=str(input_path), task_prefix=task_prefix)
        for i, item in enumerate(items, start=1)
    ]


def generate_records_from_doc(
    *,
    input_path: Path,
    input_kind: str,
    source_text: str,
    context_summary: str,
    task_prefix: str,
    max_generated_cases: int,
    doc_chunk_chars: int,
    cases_per_doc_chunk: int,
) -> list[QuickRecord]:
    return asyncio.run(
        _generate_records_from_doc_async(
            input_path=input_path,
            input_kind=input_kind,
            source_text=source_text,
            context_summary=context_summary,
            task_prefix=task_prefix,
            max_generated_cases=max_generated_cases,
            doc_chunk_chars=doc_chunk_chars,
            cases_per_doc_chunk=cases_per_doc_chunk,
        )
    )


def load_records(
    input_path: Path,
    *,
    input_kind: str,
    context_summary: str,
    task_prefix: str,
    max_generated_cases: int,
    doc_chunk_chars: int,
    cases_per_doc_chunk: int,
) -> tuple[list[QuickRecord], str]:
    if not input_path.exists():
        raise FileNotFoundError(f"输入路径不存在: {input_path}")
    parsed_cases = load_case_records(input_path)
    if input_kind == "cases":
        return parsed_cases, "cases"
    source_text = "" if parsed_cases and input_kind == "auto" else read_source_document(input_path)
    detected = detect_input_kind(input_path, source_text, parsed_cases) if input_kind == "auto" else input_kind
    if detected == "cases":
        return parsed_cases, "cases"
    if detected not in {"requirement", "analysis"}:
        raise ValueError("无法判断输入文档类型，请追加 --input-kind cases|requirement|analysis")
    records = generate_records_from_doc(
        input_path=input_path,
        input_kind=detected,
        source_text=source_text,
        context_summary=context_summary,
        task_prefix=task_prefix,
        max_generated_cases=max_generated_cases,
        doc_chunk_chars=doc_chunk_chars,
        cases_per_doc_chunk=cases_per_doc_chunk,
    )
    return records, detected


def _read_small_text(path: Path, limit: int) -> str:
    try:
        text = path.read_text(encoding="utf-8", errors="ignore")
    except OSError:
        return ""
    return text[:limit]


def _summarize_graph_dir(path: Path) -> str:
    graph_path = path / "graph.json"
    if not graph_path.is_file():
        return ""
    try:
        doc = json.loads(graph_path.read_text(encoding="utf-8"))
    except Exception:
        return ""
    nodes = doc.get("nodes") or []
    edges = doc.get("edges") or []
    parts = [f"图谱目录：{path}", f"节点数：{len(nodes)}，边数：{len(edges)}"]
    for node in nodes[:8]:
        title = node.get("title") or ""
        url = node.get("url") or node.get("url_norm") or ""
        if title or url:
            parts.append(f"- 页面：{title} {url}".strip())
    for edge in edges[:12]:
        tag = edge.get("action_tag") or edge.get("action_fingerprint") or ""
        if tag:
            parts.append(f"- 历史动作：{tag}")
    return "\n".join(parts)


def _context_dir_files(path: Path) -> list[Path]:
    """Return deterministic context files from a fixed skill context directory."""
    if not path.is_dir():
        return []
    allowed = {".md", ".txt", ".yaml", ".yml", ".json", ".py", ".ts", ".tsx", ".js", ".jsx"}
    files = [
        f
        for f in path.rglob("*")
        if f.is_file()
        and f.name not in CONTEXT_IGNORE_NAMES
        and f.suffix.lower() in allowed
        and ".skill-output" not in f.parts
    ]
    return sorted(files, key=lambda p: str(p).lower())


def _select_prompt_context_files(paths: list[Path], *, explicit_text: str) -> list[Path]:
    """Select prompt_context files for first-run query injection.

    Default is intentionally conservative: only global_readme.md is always
    injected. Product-specific files are injected only when the current user
    request explicitly names the file, its stem, or its full path.
    """
    if not paths:
        return []
    text = (explicit_text or "").lower()
    selected: list[Path] = []
    for path in paths:
        name = path.name.lower()
        if name == DEFAULT_PROMPT_CONTEXT_FILE:
            selected.append(path)
            continue
        stem = path.stem.lower()
        full = str(path).lower()
        if name in text or stem in text or full in text:
            selected.append(path)
    return sorted(selected, key=lambda p: str(p).lower())


def build_context_summary(paths: Iterable[Path], *, max_files: int = 24, max_chars_per_file: int = 1600) -> str:
    chunks: list[str] = []
    allowed = {".md", ".txt", ".yaml", ".yml", ".json", ".py", ".ts", ".tsx", ".js", ".jsx"}
    for p in paths:
        p = p.expanduser()
        if not p.exists():
            chunks.append(f"参考资料不可用：{p}")
            continue
        if p.is_dir():
            graph_summary = _summarize_graph_dir(p)
            if graph_summary:
                chunks.append(graph_summary)
                continue
            files = [f for f in p.rglob("*") if f.is_file() and f.suffix.lower() in allowed]
            ranked = sorted(
                files,
                key=lambda x: (
                    "readme" not in x.name.lower(),
                    "test" not in x.name.lower() and "case" not in x.name.lower(),
                    len(str(x)),
                ),
            )[:max_files]
            snippets = []
            for f in ranked:
                txt = _read_small_text(f, max_chars_per_file)
                if txt:
                    snippets.append(f"## {f}\n{txt}")
            if snippets:
                chunks.append(f"参考目录：{p}\n" + "\n\n".join(snippets))
        elif p.is_file():
            if p.suffix.lower() == ".json" and p.name == "graph.json":
                chunks.append(_summarize_graph_dir(p.parent))
            else:
                txt = _read_small_text(p, max_chars_per_file * 2)
                if txt:
                    chunks.append(f"参考文件：{p}\n{txt}")
    return "\n\n".join(c for c in chunks if c).strip()


def build_execution_hints(context_summary: str, *, limit: int = 3200) -> str:
    """Create self-contained query hints without file paths or reference labels."""
    if not context_summary:
        return ""
    lines: list[str] = []
    skip_prefixes = ("参考文件：", "参考目录：", "图谱目录：")
    path_re = re.compile(r"(/Users/|/(?:[A-Za-z0-9._-]+/)+|[A-Za-z]:\\)")
    for raw in context_summary.splitlines():
        line = raw.strip()
        if not line:
            continue
        if any(line.startswith(p) for p in skip_prefixes):
            continue
        if line.startswith("## "):
            title = line[3:].strip()
            if path_re.search(title):
                continue
            lines.append(f"{title}：")
            continue
        if path_re.search(line) and len(line) < 180:
            continue
        line = re.sub(r"https?://\S+", "", line).strip()
        line = re.sub(r"\s+", " ", line)
        if line:
            lines.append(line)
        if sum(len(x) for x in lines) >= limit:
            break
    text = "\n".join(lines).strip()
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text[:limit].strip()


def _with_context(query: str, execution_hints: str) -> str:
    sections = [query]
    if execution_hints:
        sections.append(
            "业务执行提示（已由参考资料吸收整理，query/YAML 独立可复跑；不要依赖外部文件或链接）：\n"
            "以下业务提示是当前被测系统的强约束；其中包含“测试前 / 先 / 必须 / 需要先”的内容，"
            "必须作为主任务前置步骤执行，不要只当作背景信息。\n"
            f"{execution_hints[:3200]}"
        )
    sections.append(
        "通用执行约束（必须遵守；query/YAML 独立可复跑）：\n"
        f"{QUICK_AGENT_GUARDRAILS}"
    )
    return "\n\n".join(sections)


def _clean_rewritten_query(text: str) -> str:
    raw = (text or "").strip()
    raw = re.sub(r"^```(?:text|markdown|md)?\s*", "", raw, flags=re.I)
    raw = re.sub(r"\s*```$", "", raw)
    lines: list[str] = []
    forbidden_prefixes = (
        "上一次执行",
        "下面是当前页面",
        "当前页面 URL",
        "当前页面标题",
        "观察摘要",
        "观察结果",
        "参考文件：",
        "参考目录：",
    )
    path_re = re.compile(r"(/Users/|/[^ \n\t:]+/|[A-Za-z]:\\)")
    for raw_line in raw.splitlines():
        line = raw_line.strip()
        if not line:
            if lines and lines[-1]:
                lines.append("")
            continue
        if any(line.startswith(prefix) for prefix in forbidden_prefixes):
            continue
        if path_re.search(line) and len(line) < 220:
            continue
        lines.append(line)
    cleaned = "\n".join(lines).strip()
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned


def _keywords_for_focus(*texts: str) -> set[str]:
    stop = {
        "当前页面",
        "页面状态",
        "页面功能",
        "布局信息",
        "可见关键文本",
        "建议下一步",
        "当前可操作控件",
        "快速测试",
        "测试目标",
        "执行",
        "验证",
        "登录",
        "使用",
        "config",
        "done",
        "success",
        "false",
        "true",
    }
    blob = "\n".join(t for t in texts if t).lower()
    words = re.findall(r"[\u4e00-\u9fff]{2,}|[a-z0-9_/-]{2,}", blob)
    return {w for w in words if w not in stop and len(w) <= 40}


def _split_hint_units(text: str) -> list[str]:
    units: list[str] = []
    buf: list[str] = []
    for raw in (text or "").splitlines():
        line = raw.strip()
        if not line:
            if buf:
                units.append(" ".join(buf))
                buf = []
            continue
        if line.startswith(("#", "-", "*", "1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.")):
            if buf:
                units.append(" ".join(buf))
                buf = []
            units.append(line.lstrip("#-* 0123456789."))
            continue
        if len(line) <= 260:
            if buf:
                units.append(" ".join(buf))
                buf = []
            units.append(line)
            continue
        buf.append(line)
        if sum(len(x) for x in buf) > 260:
            units.append(" ".join(buf))
            buf = []
    if buf:
        units.append(" ".join(buf))
    return [u.strip() for u in units if u.strip()]


def _fallback_focus_repair_context(
    *,
    record: QuickRecord,
    observe_summary: str,
    perference_hints: str,
    limit: int = REPAIR_FOCUS_LIMIT,
) -> tuple[str, str]:
    if not perference_hints:
        return "", "perference_context 为空，未提供失败修复参考。"
    target_keywords = _keywords_for_focus(record.query)
    observe_keywords = _keywords_for_focus(observe_summary)
    action_terms = (
        "入口",
        "导航",
        "菜单",
        "点击",
        "按钮",
        "tab",
        "标签",
        "弹窗",
        "等待",
        "加载",
        "新建",
        "创建",
        "保存",
        "发布",
        "提交",
        "断言",
        "预期",
        "不可点击",
        "必须",
        "不要",
        "禁止",
        "注意",
        "失败",
        "异常",
    )
    units = _split_hint_units(perference_hints)
    scored: list[tuple[int, int, str]] = []
    for idx, unit in enumerate(units):
        lower = unit.lower()
        target_overlap = sum(1 for k in target_keywords if k and k in lower)
        observe_overlap = sum(1 for k in observe_keywords if k and k in lower)
        action_score = sum(1 for term in action_terms if term.lower() in lower)
        score = target_overlap * 4 + observe_overlap * 2 + action_score
        if target_overlap > 0 or action_score >= 2:
            scored.append((score, -idx, unit))
    scored.sort(reverse=True)
    picked: list[str] = []
    total = 0
    for score, _neg_idx, unit in scored[:10]:
        normalized = re.sub(r"\s+", " ", unit).strip()
        if not normalized or normalized in picked:
            continue
        if total + len(normalized) > limit and picked:
            break
        picked.append(f"- {normalized[:360]}")
        total += len(normalized)
    focused = "\n".join(picked).strip()
    if not focused:
        return "", (
            "未从 perference_context 中匹配到与当前目标/页面观察直接相关的信息；"
            "本轮改写不注入背景参考，避免噪声。"
        )
    return focused[:limit], (
        f"启发式筛选：候选片段 {len(units)} 条，命中 {len(scored)} 条，"
        f"进入改写器 {len(picked)} 条；按原始目标和 observe 关键词相关度排序。"
    )


async def _focus_repair_context_async(
    *,
    record: QuickRecord,
    observe_summary: str,
    perference_hints: str,
    limit: int = REPAIR_FOCUS_LIMIT,
) -> str:
    from langchain_core.messages import HumanMessage, SystemMessage

    llm = _get_doc_llm()
    system = (
        "你是 UI quick 失败修复参考筛选器。只输出最多 8 条要给下一轮 query 改写使用的重点信息，不要 markdown 标题，不要解释。\n"
        "任务：从参考资料中筛选和当前失败直接相关的信息，重点是缺失前置、错误导航、页面禁忌、可见控件文案、等待/切换策略、断言口径。\n"
        "严格排除：产品背景、架构说明、无关模块、历史泛化经验、文件路径、URL、参考文件名、与当前目标/observe 无关的内容。\n"
        "每条必须能转化为下一轮 query 的具体操作约束或业务断言；如果没有相关信息，输出空字符串。"
    )
    user = (
        f"原始测试目标：\n{record.query}\n\n"
        f"当前页面观察/失败卡点：\n{observe_summary[:3500] or '(无)'}\n\n"
        f"候选参考资料摘要：\n{perference_hints[:8000]}\n\n"
        f"请输出筛选后的修复重点，总长度不超过 {limit} 字："
    )
    msg = await llm.ainvoke([SystemMessage(content=system), HumanMessage(content=user)])
    return _clean_rewritten_query(str(getattr(msg, "content", "")))[:limit].strip()


def focus_repair_context(
    *,
    record: QuickRecord,
    observe_summary: str,
    perference_hints: str,
) -> tuple[str, str]:
    if not perference_hints:
        return "", "perference_context 为空，未注入失败修复参考。"
    try:
        focused = asyncio.run(
            _focus_repair_context_async(
                record=record,
                observe_summary=observe_summary,
                perference_hints=perference_hints,
            )
        )
    except Exception as e:
        focused = ""
        fallback, note = _fallback_focus_repair_context(
            record=record,
            observe_summary=observe_summary,
            perference_hints=perference_hints,
        )
        return fallback, f"LLM 筛选失败，使用启发式筛选。{type(e).__name__}: {e}\n{note}"
    if focused:
        return focused, (
            f"LLM 筛选成功：从 {len(perference_hints)} 字失败修复参考中压缩为 "
            f"{len(focused)} 字当前失败相关重点。"
        )
    fallback, note = _fallback_focus_repair_context(
        record=record,
        observe_summary=observe_summary,
        perference_hints=perference_hints,
    )
    return fallback, f"LLM 未返回相关重点，使用启发式兜底。\n{note}"


async def _rewrite_retry_query_async(
    *,
    record: QuickRecord,
    observe_summary: str,
    prompt_hints: str,
    perference_hints: str,
) -> str:
    from langchain_core.messages import HumanMessage, SystemMessage

    llm = _get_doc_llm()
    system = (
        "你是 UI 快速测试自愈改写器。只输出一段可直接交给 browser-use quick 执行的中文 query，不要 markdown，不要解释。\n"
        "目标：吸收当前页面观察、页面特性上下文和失败修复参考，把原始测试目标改写成可独立复跑的执行任务。\n"
        "硬性约束：\n"
        "1. query 必须自包含，保留原始测试目标、入口/登录/关键数据/断言；缺失入口 URL 时不要编造。\n"
        "2. 不要出现“上一次执行未完成”“观察摘要”“当前页面 URL”“当前页面标题”等纠错提示语。\n"
        "3. 不要写入任何本地文件路径、观察文件路径、日志路径、参考文件路径或“参考文件：...”表述。\n"
        "4. 把观察结果消化为具体操作补充，例如先返回模块、切换 Tab、等待列表加载、关闭弹窗、使用页面可见文案。\n"
        "5. 失败修复参考已经过筛选，只保留与当前失败相关的重点；如参考资料与页面观察冲突，优先当前页面观察。\n"
        "6. 禁止在 query 中使用 ${USERNAME}、${PASSWORD} 等 OmegaConf 风格占位符；"
        "登录写「使用 config/.env 账号密码」或「CDP 已登录可跳过」。\n"
        "7. 若原始目标是冒烟/前端走查，必须保留同域边界：当前页面观察若显示已跳到不同域名/不同平台，只允许记录该外部页可达性和标题/错误状态，"
        "不要把外部平台导航当成新模块继续探索；重试应优先回到原始入口或同域页面继续本域模块走查。\n"
        "8. 必须保留“文本唯一性规则”：只有用例/用户未明确给出具体内容、需要你自行决定的输入文本才追加时间戳；"
        "用户已明确给出的文本、URL、文件路径、账号密码和断言期望值必须原样使用。\n"
        "9. 下拉/格式选择：原生 select 优先 get_dropdown_options + select_dropdown_option；"
        "自定义下拉才展开后精确点选项或用键盘 ArrowDown+Enter；"
        "若观察显示连续多次点同一下拉触发器但值未变，改写 query 时必须明确要求换策略。\n"
        "10. 指定模型/模式/选项：改写 query 必须保留并前置“先确认当前显示值，必要时切换到目标值；"
        "目标值确认前禁止上传、输入、提交；提交前再次确认目标值未被页面重置”。\n"
        "11. 异步/生成类任务：若 observe 或日志显示已进入生成中/创建中/处理中/任务执行中等状态，"
        "改写 query 必须要求继续等待终态，累计至少 60 秒或直到明确成功/失败；"
        "不要回头反复修复上传、下拉、输入等前置控件，前置疑点仅列人工核对。\n"
        "12. 反循环：若 observe 显示同一动作重复 ≥2 次且页面状态未变，"
        "改写 query 须禁止再次重复同动作，改为键盘选择、补必填项、关弹窗重开或 done 说明卡点。"
    )
    user = (
        f"原始测试目标：\n{record.query}\n\n"
        f"页面特性上下文（可为空，允许进入 query）：\n{prompt_hints[:3000] or '(无)'}\n\n"
        f"当前失败相关修复重点（只供你吸收，不要直接写路径或引用来源）：\n{perference_hints[:2200] or '(无)'}\n\n"
        f"当前页面观察（只供你判断下一次如何独立执行，不要照抄 URL/标题/诊断词）：\n{observe_summary[:4000] or '(无)'}\n\n"
        "请输出改写后的独立 quick query："
    )
    msg = await llm.ainvoke([SystemMessage(content=system), HumanMessage(content=user)])
    return _clean_rewritten_query(str(getattr(msg, "content", "")))


def _fallback_rewrite_retry_query(
    *,
    record: QuickRecord,
    observe_summary: str,
    prompt_hints: str,
    perference_hints: str,
) -> str:
    hints = []
    if prompt_hints:
        hints.append("页面特性约束：" + prompt_hints[:1200])
    if perference_hints:
        hints.append("失败修复参考已吸收为约束：" + perference_hints[:1200])

    observe_notes: list[str] = []
    for raw in (observe_summary or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith(("页面状态：", "页面功能：", "布局信息：", "可见关键文本：", "建议下一步：", "当前可操作控件：")):
            observe_notes.append(line)
    if observe_notes:
        hints.append("页面执行补充：" + "；".join(observe_notes)[:1400])

    query = record.query
    if hints:
        query += "\n\n补充执行约束（已吸收页面状态和参考资料，query/YAML 独立可复跑）：\n" + "\n".join(hints)
    if _is_smoke_record(record):
        query += (
            "\n\n冒烟重试边界：若当前或后续点击进入了与原始入口不同域名/不同平台的页面，"
            "只记录外部页可达性、页面标题和是否白屏/报错，不继续探索该外部平台的导航；"
            "随后返回原始入口或上一页继续同域主模块走查。"
        )
    query += (
        "\n\n请按上述目标重新执行，必要时先通过可见导航回到目标模块、等待页面加载完成、关闭阻塞弹窗；"
        "原生 select 使用 get_dropdown_options + select_dropdown_option；自定义下拉使用精确点选项或键盘操作；"
        "同一控件操作 2 次后状态未变必须换策略，禁止第 3 次重复；"
        "最终只在完成目标且断言通过时 done(success=true)，否则说明卡点。"
    )
    return _clean_rewritten_query(query)


def rewrite_retry_query(
    *,
    record: QuickRecord,
    observe_summary: str,
    prompt_hints: str,
    perference_hints: str,
    text_timestamp: str = "",
) -> tuple[str, str]:
    focused_perference_hints, focus_note = focus_repair_context(
        record=record,
        observe_summary=observe_summary,
        perference_hints=perference_hints,
    )
    advice = (
        f"# Retry Rewrite Advice\n\n"
        f"## 原始目标\n{record.query}\n\n"
        f"## 页面特性上下文\n{prompt_hints or '(空)'}\n\n"
        f"## 失败修复参考筛选\n{focus_note}\n\n"
        f"## 进入改写器的当前失败相关重点\n{focused_perference_hints or '(空；未注入背景参考，避免噪声)'}\n\n"
        f"## 当前页面观察\n{observe_summary or '(空)'}\n\n"
        "## 改写原则\n"
        "- 输出 query 必须独立可复跑，不依赖本地文件、日志或 observe 路径。\n"
        "- 不把诊断文本或大段背景直接拼入 query，只把页面状态与筛选后的修复重点转化为具体导航、等待、切换和断言补充。\n"
    )
    try:
        rewritten = asyncio.run(
            _rewrite_retry_query_async(
                record=record,
                observe_summary=observe_summary,
                prompt_hints=prompt_hints,
                perference_hints=focused_perference_hints,
            )
        )
    except Exception as e:
        advice += f"\n## LLM 改写失败\n{type(e).__name__}: {e}\n"
        rewritten = ""
    if not rewritten:
        rewritten = _fallback_rewrite_retry_query(
            record=record,
            observe_summary=observe_summary,
            prompt_hints=prompt_hints,
            perference_hints=focused_perference_hints,
        )
        advice += "\n## 使用回退策略\nLLM 改写不可用，使用启发式压缩后的独立 query。\n"
    rewritten = _append_explicit_upload_path_rule(_append_text_timestamp_rule(rewritten, text_timestamp))
    advice += f"\n## 最终重试 Query\n{rewritten}\n"
    return sanitize_quick_query_for_cli(rewritten), advice


def _regression_task_name(case_id: str) -> str:
    slug = re.sub(r"[^\w\u4e00-\u9fff-]+", "_", case_id.strip(), flags=re.UNICODE).strip("_")
    return slug or "quick_regression"


def _format_regression_yaml(*, url: str, task_name: str, description: str, query: str) -> str:
    desc = description.replace('"', "'")
    lines = [
        "target:",
        f"  url: {url}",
        "tasks:",
        f"- name: {task_name}",
        f'  description: "{desc}"',
        "  flow:",
    ]
    body = query.strip()
    if "\n" in body or len(body) > 240:
        lines.append("  - aiQuick: |")
        for line in body.splitlines():
            lines.append(f"      {line}")
    else:
        escaped = body.replace("\\", "\\\\").replace('"', '\\"')
        lines.append(f'  - aiQuick: "{escaped}"')
    return "\n".join(lines) + "\n"


async def _synthesize_regression_query_async(
    *,
    record: QuickRecord,
    attempt_trail: list[dict[str, Any]],
    winning_attempt: int,
    prompt_hints: str,
) -> str:
    from langchain_core.messages import HumanMessage, SystemMessage

    llm = _get_doc_llm()
    trail_text = []
    for item in attempt_trail:
        trail_text.append(
            f"- 第 {item['attempt'] + 1} 轮（resume={item['resume']}，returncode={item['returncode']}）：\n{item['query']}"
        )
    system = (
        "你是 UI 快速测试回归用例合成器。只输出一段可直接交给 browser-use quick 执行的中文 query，不要 markdown，不要解释。\n"
        "背景：被测环境为 CDP 模式。前若干轮 quick 曾失败，后续通过 quick --resume 续跑成功；"
        "现在需要合成一条**全新的一次性 quick query**（不使用 --resume），在 CDP 已登录前提下从当前入口页独立完成全流程并满足原始断言。\n"
        "硬性约束：\n"
        "1. 保留原始测试目标、入口 URL（若已知）、登录说明、关键数据与可观察断言；不得削弱或删除原断言。\n"
        "2. 把各轮 attempt 中已验证有效的操作路径（导航、点击文案、输入、提交方式、等待条件）消化为线性步骤，"
        "禁止依赖上一轮留在页面上的半成品状态；必要时写明「新建会话/清空对话后再提问」。\n"
        "3. 禁止出现「续跑」「上一次执行未完成」「观察摘要」「当前页面 URL/标题」等纠错或诊断用语。\n"
        "4. 禁止写入任何本地文件路径、日志路径、observe 路径或「参考文件：...」。\n"
        "5. 禁止 ${USERNAME}、${PASSWORD} 等占位符；登录写「使用 config/.env 账号密码」或「CDP 已登录可跳过」。\n"
        "6. 不要写成「等待上一轮生成完成」这类依赖续跑现场的目标；必须能在单次 quick 中从头完成。\n"
        "7. 必须保留“文本唯一性规则”：只有用例/用户未明确给出具体内容、需要你自行决定的输入文本才追加时间戳；"
        "用户已明确给出的文本、URL、文件路径、账号密码和断言期望值必须原样使用。\n"
        "8. 原生 select 优先 get_dropdown_options + select_dropdown_option；"
        "自定义下拉须精确点选项或键盘操作；同一控件 2 次无效须换策略。\n"
        "9. 指定模型/模式/选项：合成回归 query 时必须把“先确认当前显示值，必要时切换到目标值；"
        "目标值确认前禁止上传、输入、提交；提交前再次确认目标值未被页面重置”写成前置步骤。\n"
        "10. 异步/生成类任务：若提交后出现生成中/创建中/处理中/任务执行中等状态，"
        "回归 query 必须等待终态，累计至少 60 秒或直到明确成功/失败；"
        "不得把进行中的前置控件疑点写成中断条件。"
    )
    user = (
        f"原始测试目标：\n{record.query}\n\n"
        f"页面特性上下文（可为空）：\n{prompt_hints[:3000] or '(无)'}\n\n"
        f"各轮执行 query 轨迹：\n" + "\n\n".join(trail_text) + "\n\n"
        f"最终成功轮次：第 {winning_attempt + 1} 轮（CDP 续跑成功）。\n\n"
        "请输出合成后的「一次 quick 成功复现」query："
    )
    msg = await llm.ainvoke([SystemMessage(content=system), HumanMessage(content=user)])
    return _clean_rewritten_query(str(getattr(msg, "content", "")))


def _fallback_synthesize_regression_query(
    *,
    record: QuickRecord,
    attempt_trail: list[dict[str, Any]],
    prompt_hints: str,
) -> str:
    winning = attempt_trail[-1]["query"] if attempt_trail else record.query
    parts = [record.query.strip()]
    if prompt_hints:
        parts.append("页面特性：" + prompt_hints[:800])
    parts.append(
        "执行说明（CDP 一次复现，非续跑）：先确保在原始任务对应的目标模块或目标页面；"
        "如需干净会话或重置状态，请使用当前页面实际可见的等价入口；"
        "再按原始步骤完成输入、提交或其它核心操作并等待页面响应；"
        "最后严格按原始目标做可观察断言，未满足则 done(success=false)。"
    )
    if winning and winning not in record.query:
        distilled = _clean_rewritten_query(winning)
        if distilled:
            parts.append("已验证操作要点（从自愈成功轮吸收，勿照抄诊断语）：\n" + distilled[:1600])
    return sanitize_quick_query_for_cli("\n\n".join(p for p in parts if p))


def synthesize_regression_query(
    *,
    record: QuickRecord,
    attempt_trail: list[dict[str, Any]],
    winning_attempt: int,
    prompt_hints: str,
    text_timestamp: str = "",
) -> tuple[str, str]:
    notes = (
        f"# Regression Query Synthesis\n\n"
        f"## 原始目标\n{record.query}\n\n"
        f"## 成功轮次\n第 {winning_attempt + 1} 轮（CDP `--resume` 自愈成功）\n\n"
        f"## 各轮 query 轨迹\n"
    )
    for item in attempt_trail:
        notes += f"\n### 第 {item['attempt'] + 1} 轮（resume={item['resume']}，rc={item['returncode']}）\n{item['query']}\n"
    try:
        query = asyncio.run(
            _synthesize_regression_query_async(
                record=record,
                attempt_trail=attempt_trail,
                winning_attempt=winning_attempt,
                prompt_hints=prompt_hints,
            )
        )
    except Exception as e:
        notes += f"\n## LLM 合成失败\n{type(e).__name__}: {e}\n"
        query = ""
    if not query:
        notes += "\n## 使用回退策略\n"
        query = _fallback_synthesize_regression_query(
            record=record,
            attempt_trail=attempt_trail,
            prompt_hints=prompt_hints,
        )
    else:
        query = sanitize_quick_query_for_cli(query)
    query = sanitize_quick_query_for_cli(_append_text_timestamp_rule(query, text_timestamp))
    notes += f"\n## 最终回归 Query（一次 quick，非续跑）\n{query}\n"
    return query, notes


def emit_regression_archive(
    *,
    record: QuickRecord,
    log_dir: Path,
    attempt_trail: list[dict[str, Any]],
    winning_attempt: int,
    prompt_hints: str,
    text_timestamp: str,
    dry_run: bool,
) -> dict[str, str]:
    """After CDP heal-and-resume success, write one-shot regression query + YAML (no re-run)."""
    reg_dir = log_dir / REGRESSION_DIRNAME
    reg_dir.mkdir(parents=True, exist_ok=True)
    yaml_path = reg_dir / f"{record.case_id}_regression.yaml"
    query_path = reg_dir / f"{record.case_id}_regression_query.txt"
    notes_path = reg_dir / f"{record.case_id}_regression_notes.md"

    if dry_run:
        placeholder = (
            f"# {record.case_id} regression (dry-run)\n\n"
            "dry-run 未调用 LLM 合成；正式执行且 CDP 自愈成功后会写入 query/YAML。\n"
        )
        notes_path.write_text(placeholder, encoding="utf-8")
        return {
            "regression_yaml_path": str(yaml_path),
            "regression_query_path": str(query_path),
            "regression_notes_path": str(notes_path),
        }

    query, notes = synthesize_regression_query(
        record=record,
        attempt_trail=attempt_trail,
        winning_attempt=winning_attempt,
        prompt_hints=prompt_hints,
        text_timestamp=text_timestamp,
    )
    url = resolve_target_url(record.query) or extract_url(record.query) or resolve_target_url()
    if not url:
        url = extract_url(attempt_trail[0]["query"]) if attempt_trail else ""
    if not url:
        notes += "\n## 警告\n未能解析入口 URL，YAML 中 target.url 为空，归档前请手工补全。\n"

    task_name = _regression_task_name(record.case_id)
    description = (
        f"CDP 回归归档（一次 quick 复现，非续跑）；由 {record.case_id} 第 1–{winning_attempt + 1} 轮自愈成功后合成；"
        "未自动复跑验证"
    )
    yaml_text = _format_regression_yaml(
        url=url,
        task_name=task_name,
        description=description,
        query=query,
    )
    query_path.write_text(query + "\n", encoding="utf-8")
    yaml_path.write_text(yaml_text, encoding="utf-8")
    notes_path.write_text(notes, encoding="utf-8")
    return {
        "regression_query": query,
        "regression_yaml_path": str(yaml_path),
        "regression_query_path": str(query_path),
        "regression_notes_path": str(notes_path),
    }


class LiveObservationWriter:
    def __init__(self, *, cwd: Path, source_dir: Path, label: str) -> None:
        raw_root = (os.environ.get("UIWARDEN_OBSERVATION_DIR") or "").strip()
        if not raw_root:
            self.root = cwd / ".skill-output" / "uiwarden-observation"
            self.enabled = True
        else:
            root = Path(raw_root).expanduser()
            self.root = root if root.is_absolute() else cwd / root
            self.enabled = True
        safe_label = re.sub(r"[^\w.-]+", "_", label).strip("_") or "quick"
        self.target_dir = self.root / safe_label
        self.source_dir = source_dir
        self.manifest = self.target_dir / "manifest.ndjson"
        self.seen_files: dict[str, float] = {}
        self.seen_hashes: set[str] = set()
        self.index = 0

    def poll(self) -> None:
        if not self.enabled:
            return
        self.target_dir.mkdir(parents=True, exist_ok=True)
        self._copy_image_files()
        self._extract_result_json_images()

    def _copy_image_files(self) -> None:
        if not self.source_dir.exists():
            return
        for src in sorted(self.source_dir.rglob("*")):
            if not src.is_file() or src.suffix.lower() not in {".png", ".jpg", ".jpeg", ".webp"}:
                continue
            try:
                stat = src.stat()
            except OSError:
                continue
            key = str(src)
            if stat.st_size <= 0 or self.seen_files.get(key) == stat.st_mtime:
                continue
            self.seen_files[key] = stat.st_mtime
            digest = self._file_digest(src)
            if digest and digest in self.seen_hashes:
                continue
            if digest:
                self.seen_hashes.add(digest)
            self._write_frame_from_file(src, suffix=src.suffix.lower() or ".png")

    def _extract_result_json_images(self) -> None:
        if not self.source_dir.exists():
            return
        for src in sorted(self.source_dir.glob("quick_task_result*.json")):
            try:
                stat = src.stat()
            except OSError:
                continue
            key = str(src)
            if stat.st_size <= 0 or self.seen_files.get(key) == stat.st_mtime:
                continue
            self.seen_files[key] = stat.st_mtime
            try:
                raw = src.read_text(encoding="utf-8", errors="ignore")
            except OSError:
                continue
            for match in re.finditer(r'"(?:screenshot|screenshot_before_actions)"\s*:\s*"([A-Za-z0-9+/=]{2000,})"', raw):
                payload = match.group(1)
                digest = hashlib.sha1(payload[:4096].encode("ascii", errors="ignore")).hexdigest()
                if digest in self.seen_hashes:
                    continue
                self.seen_hashes.add(digest)
                try:
                    data = base64.b64decode(payload, validate=False)
                except Exception:
                    continue
                self._write_frame(data, suffix=".png", source=str(src))

    def _write_frame_from_file(self, src: Path, *, suffix: str) -> None:
        self.index += 1
        target = self.target_dir / f"frame_{self.index:04d}{suffix if suffix.startswith('.') else '.' + suffix}"
        try:
            shutil.copyfile(src, target)
        except OSError:
            return
        self._append_manifest(target, source=str(src))

    def _write_frame(self, data: bytes, *, suffix: str, source: str) -> None:
        if not data:
            return
        self.index += 1
        target = self.target_dir / f"frame_{self.index:04d}{suffix}"
        try:
            target.write_bytes(data)
        except OSError:
            return
        self._append_manifest(target, source=source)

    def _append_manifest(self, target: Path, *, source: str) -> None:
        record = {
            "framePath": str(target),
            "stepIndex": self.index,
            "actionSummary": f"browser-use 执行截图 {self.index}",
            "source": source,
            "timestamp": datetime.now().isoformat(timespec="seconds"),
        }
        with self.manifest.open("a", encoding="utf-8") as f:
            f.write(json.dumps(record, ensure_ascii=False) + "\n")

    @staticmethod
    def _file_digest(path: Path) -> str:
        try:
            return hashlib.sha1(path.read_bytes()).hexdigest()
        except OSError:
            return ""


_SOCKS_PROXY_ENV_KEYS = (
    "ALL_PROXY", "all_proxy",
    "HTTPS_PROXY", "https_proxy",
    "HTTP_PROXY", "http_proxy",
)


def _strip_socks_proxy_if_needed(env: dict[str, str]) -> None:
    """Remove SOCKS proxy env vars when socksio is not installed.

    httpx raises ImportError when it detects a SOCKS proxy (via ALL_PROXY etc.)
    but the ``socksio`` extra is missing.  Rather than crashing, we strip the
    proxy variables so the request goes direct — acceptable for browser-use
    which talks to a local Playwright browser, not through a corporate proxy.

    Additionally, httpx reads system proxy settings via
    ``urllib.request.getproxies()`` (macOS System Preferences, Windows
    Registry, or Linux env vars).  When a SOCKS proxy is detected at the
    OS level, httpx may still try to use it.  Setting ``NO_PROXY=*`` tells
    httpx to bypass all proxies.
    """
    # Check env-var-based SOCKS proxies
    has_socks_env = any(
        env.get(k, "").lower().startswith(("socks5://", "socks4://", "socks://"))
        for k in _SOCKS_PROXY_ENV_KEYS
    )
    # Check OS-level SOCKS proxy (macOS, Windows, Linux env-based)
    has_socks_system = False
    try:
        from urllib.request import getproxies
        system_proxies = getproxies()
        has_socks_system = any(
            v.lower().startswith(("socks5://", "socks4://", "socks://"))
            for v in system_proxies.values()
        )
    except Exception:
        pass

    if not has_socks_env and not has_socks_system:
        return  # No SOCKS proxy detected, nothing to do

    try:
        import socksio  # noqa: F401
        return  # socksio available, keep proxy vars
    except ImportError:
        pass

    # socksio missing — strip SOCKS env vars and set NO_PROXY=*
    if has_socks_env:
        stripped = [k for k in _SOCKS_PROXY_ENV_KEYS if env.get(k, "").lower().startswith(("socks5://", "socks4://", "socks://"))]
        for k in stripped:
            _log(f"SOCKS 代理 {k}={env[k]} 但 socksio 未安装，已清除该变量以避免 httpx ImportError")
            del env[k]
    if has_socks_system:
        _log("系统代理检测到 SOCKS 配置但 socksio 未安装，已设置 NO_PROXY=* 绕过代理")
        env["NO_PROXY"] = "*"


def _run(
    cmd: list[str],
    *,
    cwd: Path,
    dry_run: bool,
    observation_source_dir: Path | None = None,
    observation_label: str = "",
) -> int:
    print("+ " + shlex.join(cmd))
    if dry_run:
        return 0
    env = os.environ.copy()
    # Strip SOCKS proxy env vars when socksio is missing to prevent
    # httpx ImportError: "Using SOCKS proxy, but 'socksio' is not installed"
    _strip_socks_proxy_if_needed(env)
    # sitecustomize.py (CDP context-reuse monkeypatch + live screenshot capture)
    # lives in _SCRIPTS_DIR; Python auto-imports it only when that dir is on
    # PYTHONPATH. Set it unconditionally so every browser-use subprocess picks up
    # the patches — not just observation-tracked runs — otherwise CDP tasks hit
    # "Browser context management is not supported".
    scripts_dir = str(_SCRIPTS_DIR)
    existing_pythonpath = env.get("PYTHONPATH", "")
    env["PYTHONPATH"] = scripts_dir if not existing_pythonpath else scripts_dir + os.pathsep + existing_pythonpath
    if observation_source_dir is not None:
        env["UIWARDEN_OBSERVATION_LABEL"] = observation_label or observation_source_dir.name
    if observation_source_dir is None:
        return subprocess.run(cmd, cwd=str(cwd), env=env).returncode
    writer = LiveObservationWriter(
        cwd=cwd,
        source_dir=observation_source_dir,
        label=observation_label or observation_source_dir.name,
    )
    proc = subprocess.Popen(cmd, cwd=str(cwd), env=env)
    try:
        while proc.poll() is None:
            writer.poll()
            time.sleep(0.7)
        writer.poll()
        return int(proc.returncode or 0)
    finally:
        writer.poll()


def resolve_browser_use_python(requested_python: str | None = None) -> str:
    requested = (requested_python or "").strip()
    if requested:
        return os.path.expanduser(requested)
    exe, failures, from_env = find_usable_interpreter()
    if exe:
        if from_env:
            _log(f"browser-use Python 使用 UI_TEST_PYTHON/UI_TEST_ASSISTANT_PYTHON：{exe}")
        elif _exe_real_key(exe) != _exe_real_key(sys.executable):
            _log(f"自动选用可用 browser-use Python：{exe}（当前脚本解释器：{sys.executable}）")
        else:
            _log(f"browser-use Python 使用当前解释器：{exe}")
        return exe
    details = "\n\n".join(failures[-8:])
    raise RuntimeError(
        f"未找到同时满足 Python {python_range_label()}、browser-use 3.1.6 且可 import browser_use 的解释器。\n"
        f"{details}\n\n"
        "可显式设置 UI_TEST_PYTHON=/path/to/python，或使用 --python /path/to/python。"
    )


def _needs_no_zygote() -> bool:
    """Determine if Chromium needs ``--no-zygote`` to launch successfully.

    Returns True when running in a constrained environment where Chromium's
    zygote process cannot register its Mach port (macOS sandbox) or where
    the sandbox/namespace setup prevents the zygote from forking (Docker,
    Kubernetes, CI).

    Also returns True for headless Docker/CI environments where
    ``--disable-gpu`` and ``--disable-dev-shm-usage`` are recommended.
    """
    # 1. Claude Code sandbox (macOS / Linux)
    if os.environ.get("CLAUDE_WORK_DIR"):
        return True
    # Claude Code sandbox uses /tmp/claude-* for temp dirs
    import tempfile
    tmp = tempfile.gettempdir()
    if "/claude-" in tmp:
        return True
    # macOS App Sandbox
    for key in ("APP_SANDBOX_CONTAINER_ID", "APP_SANDBOX_USER_DIR"):
        if os.environ.get(key):
            return True
    # 2. Docker container detection
    if Path("/.dockerenv").is_file():
        return True
    # Docker cgroup indicator
    try:
        cgroup = Path("/proc/1/cgroup").read_text(encoding="utf-8", errors="ignore")
        if "docker" in cgroup or "kubepods" in cgroup:
            return True
    except Exception:
        pass
    # 3. CI environments
    for key in ("CI", "JENKINS_HOME", "GITLAB_CI", "GITHUB_ACTIONS", "BUILDKITE"):
        if os.environ.get(key):
            return True
    # 4. Kubernetes pod
    if os.environ.get("KUBERNETES_SERVICE_HOST"):
        return True
    return False


def _quick_cmd(
    py: str,
    query: str,
    log_dir: Path,
    *,
    no_cdp: bool,
    headless: bool,
    resume: bool,
    max_steps: int,
) -> list[str]:
    cmd = [py, "-m", "browser_use.cli", "quick"]
    if no_cdp:
        cmd.append("--no-cdp")
    # In constrained environments, force headless — there is no display anyway,
    # and headed mode causes Mach port / ProcessSingleton crashes.
    if not headless and _needs_no_zygote():
        _log("受限环境检测到，强制切换为无头模式（headed 模式在沙箱/Docker/CI 中不可用）")
        headless = True
    if headless:
        cmd.extend(["--headless", "true"])
    else:
        cmd.extend(["--headless", "false"])
    if resume:
        cmd.append("--resume")
    cmd.extend(["--max-steps", str(max_steps), query, "--log-dir", str(log_dir)])
    # In constrained environments (macOS sandbox, Docker, CI, K8s), Chromium
    # needs extra flags to launch: --no-zygote prevents Mach port / fork
    # failures; --disable-gpu avoids GPU init errors in headless servers;
    # --disable-dev-shm-usage avoids /dev/shm size issues in containers.
    if headless and _needs_no_zygote():
        cmd.append("browser.extra_chromium_args=[--no-zygote,--disable-gpu,--disable-dev-shm-usage,--disable-features=MachPortRendezvous]")
    if not headless:
        # browser-use CLI may need the inner browser override in addition to the
        # quick-level flag; keep it last so OmegaConf-style overrides are parsed.
        cmd.extend(["browser.headless", "false"])
    return cmd


def _load_records_jsonl(path: Path) -> list[QuickRecord]:
    records: list[QuickRecord] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line:
            continue
        records.append(QuickRecord(**json.loads(line)))
    return records


def _observe_cmd(py: str, output: Path, *, no_cdp: bool, log_dir: Path) -> list[str]:
    cmd = [py, "-m", "browser_use.cli", "observe"]
    if no_cdp:
        cmd.extend(["--no-cdp", "--log-dir", str(log_dir)])
    cmd.extend(["--output", str(output)])
    return cmd


def _read_observe_summary(path: Path) -> str:
    try:
        doc = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return ""
    parts: list[str] = []
    url = _clean_cell(doc.get("url"))
    title = _clean_cell(doc.get("title"))
    if url:
        parts.append(f"当前页面 URL：{url}")
    if title:
        parts.append(f"当前页面标题：{title}")
    page = doc.get("page_understanding") if isinstance(doc.get("page_understanding"), dict) else {}
    for key, label in (
        ("page_description", "页面状态"),
        ("functionality", "页面功能"),
        ("layout", "布局信息"),
    ):
        value = _clean_cell(page.get(key))
        if value:
            parts.append(f"{label}：{value[:500]}")
    for key, label in (
        ("key_visible_texts", "可见关键文本"),
        ("suggested_next_steps", "建议下一步"),
    ):
        values = page.get(key)
        if isinstance(values, list):
            clean_values = [_clean_cell(v) for v in values if _clean_cell(v)]
            if clean_values:
                parts.append(f"{label}：" + "；".join(clean_values[:10]))
    controls = doc.get("controls")
    if isinstance(controls, list) and controls:
        control_lines = []
        for c in controls[:20]:
            if not isinstance(c, dict):
                continue
            text = _clean_cell(c.get("label") or c.get("text") or c.get("dom_description"))
            purpose = _clean_cell(c.get("purpose"))
            if text or purpose:
                control_lines.append(f"{text or '未命名控件'}({purpose})")
        if control_lines:
            parts.append("当前可操作控件：" + "；".join(control_lines))
    return "\n".join(parts).strip()


def _is_effectively_empty_dir(path: Path) -> bool:
    try:
        entries = list(path.iterdir())
    except OSError:
        return False
    if not entries:
        return True
    return all(e.is_file() and e.name == ".DS_Store" for e in entries)


def cleanup_empty_log_dirs(logs_root: Path) -> None:
    if not logs_root.exists():
        return
    dirs = sorted((p for p in logs_root.rglob("*") if p.is_dir()), key=lambda p: len(p.parts), reverse=True)
    for d in dirs:
        if _is_effectively_empty_dir(d):
            for f in d.iterdir():
                if f.is_file() and f.name == ".DS_Store":
                    try:
                        f.unlink()
                    except OSError:
                        pass
            try:
                d.rmdir()
            except OSError:
                pass


def run_record(
    record: QuickRecord,
    *,
    py: str,
    out_dir: Path,
    prompt_hints: str,
    perference_hints: str,
    no_cdp: bool,
    headless: bool,
    max_retries: int,
    max_steps: int,
    text_timestamp: str,
    dry_run: bool,
    cwd: Path,
    continue_current_page: bool = False,
) -> dict[str, Any]:
    if not record.should_run:
        _log(f"跳过用例 {record.case_id} {record.title}：{record.skip_reason}")
        return {
            "case_id": record.case_id,
            "title": record.title,
            "status": "skipped",
            "reason": record.skip_reason,
            "original_case_path": str(out_dir / "logs" / record.case_id / "original_case.md"),
        }
    log_dir = out_dir / "logs" / record.case_id
    observe_dir = out_dir / "observe"
    rewrite_dir = out_dir / "rewrite"
    log_dir.mkdir(parents=True, exist_ok=True)
    observe_dir.mkdir(parents=True, exist_ok=True)
    rewrite_dir.mkdir(parents=True, exist_ok=True)
    base_query = _append_explicit_upload_path_rule(
        _append_upload_fixture_hint(
            sanitize_quick_query_for_cli(record.query),
            skill_root=_skill_root(),
        )
    )
    query = _with_context(base_query, prompt_hints)
    status = "failed"
    attempts: list[dict[str, Any]] = []
    attempt_trail: list[dict[str, Any]] = []
    regression_info: dict[str, str] = {}
    smoke_early_stop_continuation_used = False
    attempt_limit = max_retries + 1 + (1 if _is_smoke_record(record) else 0)
    _log(f"开始用例 {record.case_id}：{record.title}")
    _log(f"日志目录：{log_dir}")
    _log(f"原始用例：{log_dir / 'original_case.md'}")
    for attempt in range(attempt_limit):
        attempt_log_dir = log_dir / f"attempt_{attempt}"
        attempt_no = attempt + 1
        resume = (continue_current_page or attempt > 0) and not no_cdp
        attempt_log_dir.mkdir(parents=True, exist_ok=True)
        attempt_query_path = attempt_log_dir / "attempt_query.txt"
        attempt_query_path.write_text(query + "\n", encoding="utf-8")
        _log(
            f"第 {attempt_no} 轮执行开始（resume={resume}, max_steps={max_steps}, "
            f"no_cdp={no_cdp}, headless={headless}）"
        )
        _log_query_block(
            case_id=record.case_id,
            attempt_no=attempt_no,
            query=query,
            query_path=attempt_query_path,
        )
        rc = _run(
            _quick_cmd(
                py,
                query,
                attempt_log_dir,
                no_cdp=no_cdp,
                headless=headless,
                resume=resume,
                max_steps=max_steps,
            ),
            cwd=cwd,
            dry_run=dry_run,
            observation_source_dir=attempt_log_dir,
            observation_label=f"{record.case_id}_attempt_{attempt}",
        )
        attempt_trail.append(
            {"attempt": attempt, "returncode": rc, "query": query, "resume": resume}
        )
        result_summary, artifacts = _quick_attempt_result_summary(
            attempt_log_dir,
            returncode=rc,
            dry_run=dry_run,
        )
        _log(
            f"第 {attempt_no} 轮 quick 结果：{result_summary}\n"
            f"  result_json: {artifacts.get('result_json') or '(未生成)'}\n"
            f"  report_html: {artifacts.get('report_html') or '(未生成)'}\n"
            f"  task_yaml: {artifacts.get('task_yaml') or '(未生成)'}"
        )
        validation = validate_quick_attempt(
            record,
            attempt_log_dir,
            returncode=rc,
            dry_run=dry_run,
        )
        if _is_smoke_record(record):
            validation = validate_smoke_walkthrough_attempt(
                base_validation=validation,
                attempt_log_dir=attempt_log_dir,
                returncode=rc,
                dry_run=dry_run,
            )
        elif _is_health_record(record):
            validation = validate_health_check_attempt(
                base_validation=validation,
                attempt_log_dir=attempt_log_dir,
                returncode=rc,
                dry_run=dry_run,
            )
        validation_path = write_validation_artifact(attempt_log_dir, validation)
        attempts.append(
            {
                "attempt": attempt,
                "returncode": rc,
                "query": query,
                "resume": resume,
                "query_path": str(attempt_query_path),
                "result_summary": result_summary,
                "validation": validation.to_dict(),
                "validation_path": str(validation_path),
                **artifacts,
            }
        )
        _log(
            f"第 {attempt_no} 轮完成度验收：{'通过' if validation.passed else '未通过'} "
            f"（{validation.method}）— {validation.reason}"
        )
        if rc == 0 and validation.passed:
            status = "success"
            _log(f"第 {attempt_no} 轮执行成功（CLI + 验收）：{record.case_id}")
            if not no_cdp and attempt > 0:
                _log(
                    f"CDP 自愈续跑已成功（第 1–{attempt_no} 轮），开始合成一次复现回归 query/YAML（不自动复跑）"
                )
                regression_info = emit_regression_archive(
                    record=record,
                    log_dir=log_dir,
                    attempt_trail=attempt_trail,
                    winning_attempt=attempt,
                    prompt_hints=prompt_hints,
                    text_timestamp=text_timestamp,
                    dry_run=dry_run,
                )
                attempts[-1]["regression_yaml_path"] = regression_info.get("regression_yaml_path", "")
                attempts[-1]["regression_query_path"] = regression_info.get("regression_query_path", "")
                attempts[-1]["regression_notes_path"] = regression_info.get("regression_notes_path", "")
                _log(
                    "回归归档已写入（供后续反复执行，本轮未验证复跑）：\n"
                    f"  YAML: {regression_info.get('regression_yaml_path', '')}\n"
                    f"  Query: {regression_info.get('regression_query_path', '')}\n"
                    f"  说明: {regression_info.get('regression_notes_path', '')}"
                )
                _log("如需验证归档可用性，请明确要求后再对该 YAML/query 执行 quick；默认不在自愈成功后自动复跑。")
            break
        if rc == 0 and not validation.passed:
            _log(
                f"第 {attempt_no} 轮 CLI 返回成功但验收未通过（假阳性已拦截），"
                f"日志：{attempt_log_dir}，验收：{validation_path}"
            )
        else:
            _log(f"第 {attempt_no} 轮执行失败：returncode={rc}，本轮日志：{attempt_log_dir}")
        if (
            _is_smoke_record(record)
            and not smoke_early_stop_continuation_used
            and _is_smoke_early_stop_validation(validation)
        ):
            smoke_early_stop_continuation_used = True
            query = _smoke_early_stop_continuation_query(query, validation)
            attempts[-1]["smoke_early_stop_continuation"] = True
            _log(
                "检测到冒烟任务在中间状态提前 done，追加 1 轮 browser-use 纠偏续跑；"
                "这轮会生成独立 quick 执行报告。"
            )
            continue
        if attempt >= max_retries:
            _log(f"已达到最大重试次数 {max_retries}，停止该用例：{record.case_id}")
            break
        obs = observe_dir / f"{record.case_id}_observe_{attempt + 1}.json"
        _log(f"开始 observe 当前页面：{obs}")
        _run(_observe_cmd(py, obs, no_cdp=no_cdp, log_dir=attempt_log_dir), cwd=cwd, dry_run=dry_run)
        observe_summary = "" if dry_run else _read_observe_summary(obs)
        if observe_summary:
            _log("observe 摘要：\n" + _clip_for_log(observe_summary, limit=LOG_SUMMARY_LIMIT))
        else:
            _log("observe 摘要为空，将基于原始目标和参考信息改写。")
        _log(f"开始改写为第 {attempt_no + 1} 轮 query")
        query, advice = rewrite_retry_query(
            record=record,
            observe_summary=observe_summary,
            prompt_hints=prompt_hints,
            perference_hints=perference_hints,
            text_timestamp=text_timestamp,
        )
        advice_path = rewrite_dir / f"{record.case_id}_attempt_{attempt + 1}_advice.md"
        query_path = rewrite_dir / f"{record.case_id}_attempt_{attempt + 1}_query.txt"
        advice_path.write_text(advice, encoding="utf-8")
        query_path.write_text(query + "\n", encoding="utf-8")
        attempts[-1]["rewrite_advice_path"] = str(advice_path)
        attempts[-1]["rewrite_query_path"] = str(query_path)
        _log(f"改写建议已写入：{advice_path}")
        _log(f"第 {attempt_no + 1} 轮将使用 query：\n{_clip_for_log(query)}")
    _log(f"用例完成：{record.case_id} status={status}")
    result: dict[str, Any] = {
        "case_id": record.case_id,
        "title": record.title,
        "status": status,
        "attempts": attempts,
        "original_case_path": str(log_dir / "original_case.md"),
    }
    if regression_info:
        result["regression"] = regression_info
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="ui-test-assistant quick 单任务/批量执行辅助入口")
    parser.add_argument("--query", default="", help="单条 quick 自然语言任务")
    parser.add_argument(
        "--smoke-walkthrough",
        action="store_true",
        help="将 --query 解释为冒烟验证/前端走查意图，合成可见一级主模块可用性检查 prompt，并输出冒烟走查报告",
    )
    parser.add_argument(
        "--health-check",
        action="store_true",
        help="将 --query 解释为页面巡检/探活意图，只检查入口页可访问、渲染和基本加载，并输出页面巡检报告",
    )
    parser.add_argument("--input", type=Path, default=None, help="批量输入：测试用例集合、需求文档或测分文档")
    parser.add_argument(
        "--input-kind",
        choices=INPUT_KINDS,
        default="auto",
        help="输入类型：auto 自动识别；cases=用例集合；requirement=需求文档；analysis=测分/测试分析文档",
    )
    parser.add_argument(
        "--target-url",
        default="",
        help="输入文档未写入口 URL 时，用该 URL 作为每条 quick query 的目标入口",
    )
    parser.add_argument(
        "--task-prefix",
        default="",
        help="附加到每条 quick query 前的公共说明，例如入口 URL、登录态、租户、执行范围",
    )
    parser.add_argument(
        "--user-query",
        default="",
        help="用户原始请求，仅用于配置来源校验与产物追溯；不会直接作为 quick 任务执行",
    )
    parser.add_argument(
        "--max-generated-cases",
        type=int,
        default=0,
        help="需求/测分文档拆解出的最大用例数；默认 0 表示不限制，避免多需求文档只生成少量代表性用例",
    )
    parser.add_argument(
        "--doc-chunk-chars",
        type=int,
        default=DEFAULT_DOC_CHUNK_CHARS,
        help=f"需求/测分文档按章节合并后的单片段最大字符数，默认 {DEFAULT_DOC_CHUNK_CHARS}",
    )
    parser.add_argument(
        "--cases-per-doc-chunk",
        type=int,
        default=DEFAULT_CASES_PER_DOC_CHUNK,
        help=f"每个文档片段最多生成的用例数；0 表示不限制，默认 {DEFAULT_CASES_PER_DOC_CHUNK}",
    )
    parser.add_argument(
        "--context",
        action="append",
        type=Path,
        default=[],
        help=(
            "显式辅助参考路径，可重复；仅在 config/perference_context 为空时作为失败自愈参考，"
            "不会直接拼入 quick query"
        ),
    )
    parser.add_argument("--output-dir", type=Path, required=True, help="quick 产物目录")
    parser.add_argument(
        "--cdp",
        action="store_true",
        help="使用本地 Chrome CDP/本地登录态。默认不使用 CDP，除非用户明确要求本地登录态、调试、本地浏览器或 CDP。",
    )
    parser.add_argument(
        "--no-cdp",
        action="store_true",
        help="显式使用 --no-cdp（新启浏览器，非附着 CDP）。这是 quick runner 默认行为，保留该参数用于兼容旧命令。",
    )
    parser.add_argument(
        "--headless",
        action="store_true",
        help="显式无头模式。quick runner 默认无头，保留该参数用于兼容旧命令。",
    )
    parser.add_argument(
        "--headed",
        action="store_true",
        help="有头/可视化模式：向 browser-use quick 传 --headless false，并追加 browser.headless false。",
    )
    parser.add_argument(
        "--records-jsonl",
        type=Path,
        default=None,
        help="复用已有 batch/records.jsonl，跳过需求/用例拆解（与 --input 二选一）",
    )
    parser.add_argument(
        "--max-retries",
        type=int,
        default=None,
        help="每条失败后 observe 改写重试次数；未指定时读 config/.env 的 QUICK_MAX_RETRIES（默认 1）",
    )
    parser.add_argument(
        "--max-steps",
        type=int,
        default=None,
        help="quick ReAct 最大步数；未指定时读 config/.env 的 QUICK_MAX_STEPS（默认 40）",
    )
    parser.add_argument("--dry-run", action="store_true", help="只生成 records/queries/context，不调用 browser-use")
    parser.add_argument(
        "--no-login",
        action="store_true",
        help="跳过登录凭据预检（适用于无需登录的公开页面，或已通过 CDP 使用本地登录态的场景）",
    )
    parser.add_argument(
        "--username",
        default="",
        help="登录账号（AI Agent 可从对话上下文中传入，无需手动编辑 .env；写入 USERNAME 等环境变量供 browser-use 使用）",
    )
    parser.add_argument(
        "--password",
        default="",
        help="登录密码（AI Agent 可从对话上下文中传入，无需手动编辑 .env；写入 PASSWORD 等环境变量供 browser-use 使用）",
    )
    parser.add_argument(
        "--skip-config-check",
        action="store_true",
        help="跳过 URL/凭据/上下文路径预检（仅调试）",
    )
    parser.add_argument(
        "--python",
        default="",
        help=(
            "执行 browser-use 的 Python。未指定时自动探测 UI_TEST_PYTHON/UI_TEST_ASSISTANT_PYTHON、"
            "当前解释器和 PATH 上的 python3.x。"
        ),
    )
    parser.add_argument("--skill-root", type=Path, default=None, help="技能根目录")
    args = parser.parse_args()

    if not args.query and args.input is None and args.records_jsonl is None:
        print("必须提供 --query、--input 或 --records-jsonl", file=sys.stderr)
        return 2
    if args.records_jsonl is not None and args.input is not None:
        print("--records-jsonl 与 --input 不能同时使用", file=sys.stderr)
        return 2
    if args.smoke_walkthrough and args.health_check:
        print("--smoke-walkthrough 与 --health-check 不能同时使用", file=sys.stderr)
        return 2
    if args.smoke_walkthrough and (args.input is not None or args.records_jsonl is not None):
        print("--smoke-walkthrough 仅支持单条 --query，不支持 --input / --records-jsonl", file=sys.stderr)
        return 2
    if args.health_check and (args.input is not None or args.records_jsonl is not None):
        print("--health-check 仅支持单条 --query，不支持 --input / --records-jsonl", file=sys.stderr)
        return 2
    if args.cdp and args.no_cdp:
        print("--cdp 与 --no-cdp 不能同时使用", file=sys.stderr)
        return 2
    if args.headed and args.headless:
        print("--headed 与 --headless 不能同时使用", file=sys.stderr)
        return 2

    skill_root = (args.skill_root or _skill_root()).resolve()
    _load_env(skill_root)
    _log_browser_profile_selection()
    try:
        uiwarden_execution = _enforce_uiwarden_browser_execution()
    except RuntimeError as error:
        print(str(error), file=sys.stderr)
        return 2
    if uiwarden_execution and args.no_cdp:
        print("UIWarden 可控浏览器会话禁止 --no-cdp；不会回退新建 Playwright 浏览器", file=sys.stderr)
        return 2
    # Normalize CLI-provided credentials into env vars so browser-use can use
    # them for login. This lets AI agents pass credentials from conversation
    # context without requiring manual .env edits.
    cli_username = (args.username or "").strip()
    cli_password = (args.password or "").strip()
    if cli_username:
        for key in ("QUICK_NAME", "QUICK_USERNAME", "USERNAME", "user"):
            os.environ[key] = cli_username
        _log(f"登录账号来源：CLI --username（{_mask_credential(cli_username)}）")
    if cli_password:
        for key in ("QUICK_PASSWORK", "QUICK_PASSWORD", "PASSWORD", "pass"):
            os.environ[key] = cli_password
        _log("登录密码来源：CLI --password")
    try:
        args.python = resolve_browser_use_python(args.python)
    except RuntimeError as e:
        print(str(e), file=sys.stderr)
        return 2
    quick_max_steps = resolve_quick_max_steps(args.max_steps)
    if args.health_check and args.max_steps is None and quick_max_steps > 25:
        quick_max_steps = 25
    if args.smoke_walkthrough and args.max_steps is None and quick_max_steps < 70:
        quick_max_steps = 70
    if (args.smoke_walkthrough or args.health_check) and args.max_retries is None:
        quick_max_retries = 0
    else:
        quick_max_retries = resolve_quick_max_retries(args.max_retries)
    mode_intent_text = "\n".join(
        p for p in (args.user_query, args.query, args.task_prefix, args.target_url) if p
    )
    inferred_cdp = _has_intent(mode_intent_text, CDP_INTENT_PATTERNS)
    inferred_headed = _has_intent(mode_intent_text, HEADED_INTENT_PATTERNS)
    inferred_no_login = _has_intent(mode_intent_text, NO_LOGIN_INTENT_PATTERNS)
    use_no_cdp = not (bool(args.cdp) or (inferred_cdp and not args.no_cdp))
    use_headless = not (bool(args.headed) or (inferred_headed and not args.headless))
    if uiwarden_execution:
        use_no_cdp = False
        use_headless = False
    use_current_page = bool(uiwarden_execution) and not any((
        args.target_url.strip(),
        extract_url(args.query or ""),
        extract_url(args.user_query or ""),
        extract_url(args.task_prefix or ""),
    ))
    continue_current_page = use_current_page and bool(uiwarden_execution.get("continueCurrentPage"))
    auto_no_login = bool(args.no_login) or inferred_no_login or bool(uiwarden_execution)
    text_timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    _log(
        f"启动 quick runner：output_dir={args.output_dir.expanduser()}，max_retries={quick_max_retries}，"
        f"max_steps={quick_max_steps}，no_cdp={use_no_cdp}，headless={use_headless}，"
        f"no_login={auto_no_login}，dry_run={args.dry_run}，"
        f"text_timestamp={text_timestamp}"
    )

    prefix_parts = []
    if args.target_url.strip():
        prefix_parts.append(f"入口 URL：{args.target_url.strip()}")
    elif args.input is not None or args.records_jsonl is not None:
        resolved_url = resolve_target_url(
            args.query or "",
            explicit_url="",
            task_prefix=args.task_prefix.strip(),
            user_query=args.user_query.strip(),
        )
        if resolved_url:
            prefix_parts.append(f"入口 URL：{resolved_url}")
    if args.task_prefix.strip():
        prefix_parts.append(args.task_prefix.strip())
    task_prefix = "。".join(prefix_parts)

    if not args.skip_config_check:
        check = preflight_quick(
            args.query or "",
            explicit_target_url=args.target_url.strip(),
            task_prefix=task_prefix,
            user_query=args.user_query.strip(),
            context_paths=args.context,
            skill_root=skill_root,
            require_url=not bool(uiwarden_execution),
            require_login=not (args.smoke_walkthrough or args.health_check or auto_no_login),
            cli_username=cli_username,
            cli_password=cli_password,
        )
        if not check.ok:
            print(check.message, file=sys.stderr)
            return 3
        if args.query and not use_current_page:
            args.query = check.enriched_query
        elif use_current_page:
            mode = "续轮恢复" if continue_current_page else "首轮当前页"
            _log(f"UIWarden 可控浏览器{mode}：忽略 .env QUICK_URL，使用当前锁定标签页")
    smoke_focus = ""
    custom_smoke = False
    if args.smoke_walkthrough:
        smoke_url = resolve_target_url(
            args.query or "",
            explicit_url=args.target_url.strip(),
            task_prefix=task_prefix,
            user_query=args.user_query.strip(),
        )
        smoke_focus = smoke_focus_from_query(
            user_query=args.user_query.strip() or args.query,
            target_url=smoke_url,
            task_prefix=task_prefix,
        )
        custom_smoke = is_custom_smoke_focus(smoke_focus)
        args.query = build_smoke_walkthrough_query(
            user_query=args.user_query.strip() or args.query,
            target_url=smoke_url,
            task_prefix=task_prefix,
        )
    elif args.health_check:
        health_url = resolve_target_url(
            args.query or "",
            explicit_url=args.target_url.strip(),
            task_prefix=task_prefix,
            user_query=args.user_query.strip(),
        )
        args.query = build_health_check_query(
            user_query=args.user_query.strip() or args.query,
            target_url=health_url,
            task_prefix=task_prefix,
        )

    out_dir = args.output_dir.expanduser()
    batch_dir = out_dir / "batch"
    context_dir = out_dir / "context"
    batch_dir.mkdir(parents=True, exist_ok=True)
    context_dir.mkdir(parents=True, exist_ok=True)

    all_prompt_context_paths = _context_dir_files(skill_root / "config" / PROMPT_CONTEXT_DIRNAME)
    prompt_context_request_text = "\n".join(
        p
        for p in (
            args.query or "",
            args.user_query or "",
            args.task_prefix or "",
            args.target_url or "",
        )
        if p
    )
    prompt_context_paths = _select_prompt_context_files(
        all_prompt_context_paths,
        explicit_text=prompt_context_request_text,
    )
    perference_context_paths = _context_dir_files(skill_root / "config" / PERFERENCE_CONTEXT_DIRNAME)
    user_perference_paths = args.context if not perference_context_paths else []
    _log(
        "上下文来源："
        f"prompt_context={len(prompt_context_paths)}/{len(all_prompt_context_paths)} 个文件，"
        f"perference_context={len(perference_context_paths)} 个文件，"
        f"user_context_as_perference={len(user_perference_paths)} 个路径"
        + ("，quick_lite_skip_prompt_context=True" if (args.smoke_walkthrough or args.health_check) and prompt_context_paths else "")
    )

    prompt_context_summary = build_context_summary(prompt_context_paths)
    if prompt_context_summary:
        (context_dir / "prompt_context_summary.md").write_text(prompt_context_summary, encoding="utf-8")
    prompt_hints = "" if (args.smoke_walkthrough or args.health_check) else build_execution_hints(prompt_context_summary)
    if prompt_hints:
        (context_dir / "prompt_execution_hints.md").write_text(prompt_hints, encoding="utf-8")
    elif (args.smoke_walkthrough or args.health_check) and prompt_context_summary:
        (context_dir / "prompt_execution_hints.skipped.md").write_text(
            "冒烟走查/页面巡检首轮默认不注入 config/prompt_context，避免固定业务系统提示污染跨站点判断。\n",
            encoding="utf-8",
        )

    perference_context_summary = build_context_summary(perference_context_paths or user_perference_paths)
    if perference_context_summary:
        (context_dir / "perference_context_summary.md").write_text(perference_context_summary, encoding="utf-8")
    perference_hints = build_execution_hints(perference_context_summary, limit=3600)
    if perference_hints:
        (context_dir / "perference_execution_hints.md").write_text(perference_hints, encoding="utf-8")
    context_source = {
        "prompt_context_dir": str(skill_root / "config" / PROMPT_CONTEXT_DIRNAME),
        "prompt_context_files": [str(p) for p in prompt_context_paths],
        "prompt_context_all_files": [str(p) for p in all_prompt_context_paths],
        "prompt_context_selection_rule": (
            f"default_only_{DEFAULT_PROMPT_CONTEXT_FILE}; specific_files_require_explicit_name_or_path"
        ),
        "perference_context_dir": str(skill_root / "config" / PERFERENCE_CONTEXT_DIRNAME),
        "perference_context_files": [str(p) for p in perference_context_paths],
        "user_context_used_as_perference": [str(p) for p in user_perference_paths],
        "ignored_user_context_because_perference_context_not_empty": [
            str(p) for p in (args.context if perference_context_paths else [])
        ],
    }
    (context_dir / "context_sources.json").write_text(
        json.dumps(context_source, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    doc_generation_context = "\n\n".join(
        c for c in (prompt_context_summary, perference_context_summary) if c
    )

    detected_kind = "query"
    if args.records_jsonl is not None:
        records_path = args.records_jsonl.expanduser()
        _log(f"复用已有用例列表：{records_path}")
        try:
            records = _load_records_jsonl(records_path)
        except Exception as e:
            print(f"读取 records.jsonl 失败: {e}", file=sys.stderr)
            return 2
        detected_kind = "records_reuse"
        _log(f"用例列表加载完成：records={len(records)}")
    elif args.input is not None:
        _log(f"开始解析输入：{args.input.expanduser()}（input_kind={args.input_kind}）")
        try:
            records, detected_kind = load_records(
                args.input.expanduser(),
                input_kind=args.input_kind,
                context_summary=doc_generation_context,
                task_prefix=task_prefix,
                max_generated_cases=max(0, args.max_generated_cases),
                doc_chunk_chars=max(4000, args.doc_chunk_chars),
                cases_per_doc_chunk=max(0, args.cases_per_doc_chunk),
            )
        except Exception as e:
            print(f"解析输入失败: {e}", file=sys.stderr)
            return 2
        _log(f"输入解析完成：detected_kind={detected_kind}，records={len(records)}")
    else:
        if args.smoke_walkthrough:
            detected_kind = "smoke_walkthrough"
            case_id = "001_smoke_walkthrough"
            title = "冒烟验证 / 前端走查"
            raw = {
                "title": title,
                "test_point": "可见一级主模块与核心主页面前端可用性",
                "steps": [
                    "打开入口并等待主页面稳定",
                    "先视觉盘点左侧/顶部导航、可展开分组和当前页面标题，列出可见一级主模块",
                    "逐个进入可见一级主模块或代表性子入口，检查主内容、关键列表/卡片/表单/搜索区/操作按钮是否可见",
                    "每个模块轻量点击或切换一到两个只读或可撤销核心入口",
                    "观察红色错误/警告、异常 toast/modal/banner、接口失败提示、白屏、长期 loading、组件遮挡重叠等疑似 bug",
                    "按 SMOKE_REPORT 结构输出结论、模块清单、模块覆盖、已检查、疑似问题、未覆盖和建议",
                ],
                "expected": "可见一级主模块可访问、核心区域可见、轻量交互正常；若存在疑似 bug，最终报告必须列出位置、现象、复现动作和证据文本。",
                "case_type": "smoke_walkthrough",
                "custom_smoke": custom_smoke,
                "smoke_focus": smoke_focus,
            }
        elif args.health_check:
            detected_kind = "health_check"
            case_id = "001_health_check"
            title = "页面巡检 / 探活"
            raw = {
                "title": title,
                "test_point": "入口页面可访问、正常渲染和完成基本加载",
                "steps": [
                    "打开入口并等待页面稳定",
                    "确认首屏主体、标题、导航、登录表单或错误页状态",
                    "观察白屏、崩溃、404/500、长期 loading、明显红色错误或资源加载失败",
                    "按 HEALTH_REPORT 结构输出结论、入口、最终页面、渲染状态、异常和备注",
                ],
                "expected": "页面可访问并正常渲染；若存在探活异常，最终报告必须列出严重级别、现象和证据文本。",
                "case_type": "health_check",
            }
        else:
            detected_kind = "query"
            case_id = "001_quick"
            title = args.query[:80] or "quick"
            raw = None
        records = [
            QuickRecord(
                case_id=case_id,
                title=title,
                query="。".join(p for p in (task_prefix, args.query) if p),
                source="query",
                raw=raw,
            )
        ]
        if args.smoke_walkthrough:
            _log("冒烟走查 quick 模式：records=1")
        elif args.health_check:
            _log("页面巡检 quick 模式：records=1")
        else:
            _log("单任务 quick 模式：records=1")
    if not records:
        print("未解析到可执行测试记录", file=sys.stderr)
        return 2
    records = [
        normalize_record_query_for_execution(
            record,
            task_prefix=task_prefix,
            text_timestamp=text_timestamp,
        )
        for record in records
    ]

    (batch_dir / "records.jsonl").write_text(
        "\n".join(json.dumps(asdict(r), ensure_ascii=False) for r in records) + "\n",
        encoding="utf-8",
    )
    (batch_dir / "input_kind.txt").write_text(detected_kind + "\n", encoding="utf-8")
    (batch_dir / "queries.md").write_text(
        "\n\n".join(f"## {r.case_id} {r.title}\n\n{r.query}" for r in records),
        encoding="utf-8",
    )
    write_original_case_artifacts(records, batch_dir=batch_dir, out_dir=out_dir)

    runnable = sum(1 for r in records if r.should_run)
    skipped = len(records) - runnable
    _log(f"准备执行：总记录={len(records)}，可执行={runnable}，跳过={skipped}，records={batch_dir / 'records.jsonl'}")
    results = [
        run_record(
            r,
            py=args.python,
            out_dir=out_dir,
            prompt_hints=prompt_hints,
            perference_hints=perference_hints,
            no_cdp=use_no_cdp,
            headless=use_headless,
            max_retries=quick_max_retries,
            max_steps=quick_max_steps,
            text_timestamp=text_timestamp,
            dry_run=args.dry_run,
            cwd=Path.cwd(),
            continue_current_page=continue_current_page,
        )
        for r in records
    ]
    cleanup_empty_log_dirs(out_dir / "logs")
    summary_lines = [
        "| case_id | title | status | validation | original_case |",
        "| --- | --- | --- | --- | --- |",
    ]
    for item in results:
        original_case_path = str(item.get("original_case_path") or "").replace("|", "/")
        val_note = ""
        attempts_list = item.get("attempts") or []
        if attempts_list:
            last_val = (attempts_list[-1].get("validation") or {})
            if last_val:
                val_note = _clean_cell(last_val.get("reason", ""))[:80]
        summary_lines.append(
            f"| {item['case_id']} | {str(item['title']).replace('|', '/')} | {item['status']} | "
            f"{val_note.replace('|', '/')} | {original_case_path} |"
        )
    (batch_dir / "results.md").write_text("\n".join(summary_lines) + "\n", encoding="utf-8")
    (batch_dir / "results.json").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    smoke_payload: dict[str, Any] | None = None
    if args.smoke_walkthrough:
        smoke_payload = write_smoke_walkthrough_report(records=records, results=results, batch_dir=batch_dir)
        if smoke_payload["issues"]:
            _log("冒烟走查疑似问题：\n" + "\n".join(f"- {issue}" for issue in smoke_payload["issues"]))
        else:
            _log("冒烟走查疑似问题：无")
        ai_analysis = smoke_payload.get("ai_analysis") or {}
        if ai_analysis:
            _log(
                "冒烟 AI 分析结论："
                f"{ai_analysis.get('verdict', '未知')} — {ai_analysis.get('reason', '')}"
            )
            coverage = ai_analysis.get("coverage_points") or []
            problems = ai_analysis.get("problem_points") or []
            manual = ai_analysis.get("manual_review_points") or []
            _log("冒烟 AI 覆盖点：\n" + "\n".join(f"- {item}" for item in coverage[:8]))
            _log("冒烟 AI 问题点：\n" + "\n".join(f"- {item}" for item in problems[:8]))
            _log("冒烟 AI 需人工核对：\n" + "\n".join(f"- {item}" for item in manual[:8]))
        _log(f"冒烟走查报告：{batch_dir / 'smoke_walkthrough_report.md'}")
        _log(f"冒烟 AI 分析：{batch_dir / 'smoke_ai_analysis.md'}")
    health_payload: dict[str, Any] | None = None
    if args.health_check:
        health_payload = write_health_check_report(records=records, results=results, batch_dir=batch_dir)
        if health_payload["issues"]:
            _log("页面巡检异常：\n" + "\n".join(f"- {issue}" for issue in health_payload["issues"]))
        else:
            _log("页面巡检异常：无")
        _log(f"页面巡检报告：{batch_dir / 'health_check_report.md'}")
    ok = all(r["status"] in ("success", "skipped") for r in results)
    _log(f"执行汇总：success={sum(1 for r in results if r['status'] == 'success')}，skipped={sum(1 for r in results if r['status'] == 'skipped')}，failed={sum(1 for r in results if r['status'] == 'failed')}，results={batch_dir / 'results.md'}")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
